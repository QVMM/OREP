"""竞赛助手：本地文档文本抽取（PDF / DOCX / PPTX / 纯文本 + 扫描件 OCR）。"""
from __future__ import annotations

import logging
import os
import re
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_EXTS = {
    "txt", "md", "markdown", "csv", "json", "xml", "html", "htm",
    "java", "py", "js", "ts", "vue", "css", "scss", "sql", "yml", "yaml", "log",
}


def _env_bool(name: str, default: bool = True) -> bool:
    raw = os.getenv(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def _ocr_enabled() -> bool:
    return _env_bool("ASSISTANT_OCR_ENABLED", True)


def _ocr_lang() -> str:
    return os.getenv("ASSISTANT_OCR_LANG", "chi_sim+eng")


def _ocr_max_pages() -> int:
    try:
        return max(1, min(30, int(os.getenv("ASSISTANT_OCR_MAX_PAGES", "8"))))
    except Exception:
        return 8


def _ocr_min_text_chars() -> int:
    """文字层少于此长度时，视为可能扫描件并尝试 OCR。"""
    try:
        return max(0, int(os.getenv("ASSISTANT_OCR_MIN_TEXT_CHARS", "40")))
    except Exception:
        return 40


def _ocr_zoom() -> float:
    try:
        return max(1.0, min(3.0, float(os.getenv("ASSISTANT_OCR_ZOOM", "2.0"))))
    except Exception:
        return 2.0


def _truncate(text: str, max_chars: int = 14000) -> str:
    text = (text or "").strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n…(已截断)"


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_file(path: Path, max_chars: int = 14000) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return _truncate(_clean(raw.decode(enc)), max_chars)
        except UnicodeDecodeError:
            continue
    return _truncate(_clean(raw.decode("utf-8", errors="ignore")), max_chars)


def _extract_pdf_text_layer(path: Path, max_chars: int = 14000) -> str:
    """仅抽取 PDF 文字层（不做 OCR）。"""
    text = ""
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(str(path))
        parts: list[str] = []
        total = 0
        for i, page in enumerate(doc):
            try:
                t = page.get_text("text") or ""
            except Exception:
                t = ""
            if t.strip():
                chunk = f"\n--- 第 {i + 1} 页 ---\n{t.strip()}"
                parts.append(chunk)
                total += len(chunk)
                if total >= max_chars:
                    break
        doc.close()
        text = _clean("\n".join(parts))
    except Exception as e:
        logger.info("fitz extract failed, fallback PyPDF2: %s", e)

    if text:
        return _truncate(text, max_chars)

    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(str(path))
        parts = []
        total = 0
        for i, page in enumerate(reader.pages):
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                chunk = f"\n--- 第 {i + 1} 页 ---\n{t.strip()}"
                parts.append(chunk)
                total += len(chunk)
                if total >= max_chars:
                    break
        text = _clean("\n".join(parts))
    except Exception as e:
        raise RuntimeError(f"PDF 解析失败: {e}") from e

    return _truncate(text, max_chars) if text else ""


def ocr_pdf_pages(path: Path, max_chars: int = 14000, on_progress=None) -> str:
    """
    扫描件 OCR：PyMuPDF 渲染页面为图片 → Tesseract。
    on_progress(page:int, total:int, title:str|None) 可选回调。
    """
    if not _ocr_enabled():
        return ""

    try:
        import fitz
        import pytesseract
        from PIL import Image
    except Exception as e:
        logger.warning("OCR dependencies missing: %s", e)
        return ""

    tess_cmd = os.getenv("TESSERACT_CMD") or "/opt/homebrew/bin/tesseract"
    if Path(tess_cmd).exists():
        pytesseract.pytesseract.tesseract_cmd = tess_cmd
    tessdata = os.getenv("TESSDATA_PREFIX") or "/opt/homebrew/share/tessdata"
    if Path(tessdata).exists():
        os.environ["TESSDATA_PREFIX"] = tessdata

    lang = _ocr_lang()
    max_pages = _ocr_max_pages()
    zoom = _ocr_zoom()
    mat = fitz.Matrix(zoom, zoom)

    doc = fitz.open(str(path))
    parts: list[str] = []
    total = 0
    full_n = len(doc)
    page_count = min(full_n, max_pages)
    try:
        for i in range(page_count):
            if on_progress:
                try:
                    on_progress(i + 1, page_count, path.name)
                except Exception:
                    pass
            page = doc[i]
            try:
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                gray = img.convert("L")
                t = pytesseract.image_to_string(gray, lang=lang) or ""
            except Exception as e:
                logger.warning("OCR page %s failed: %s", i + 1, e)
                t = ""
            if t.strip():
                chunk = f"\n--- 第 {i + 1} 页（OCR） ---\n{t.strip()}"
                parts.append(chunk)
                total += len(chunk)
                if total >= max_chars:
                    break
    finally:
        doc.close()

    text = _clean("\n".join(parts))
    if not text:
        return ""
    note = ""
    if full_n > page_count:
        note = f"\n\n（OCR 仅处理前 {page_count}/{full_n} 页，可提高 ASSISTANT_OCR_MAX_PAGES）"
    return _truncate(text + note, max_chars + 200)


def extract_pdf(path: Path, max_chars: int = 14000, on_progress=None) -> tuple[str, str]:
    """
    返回 (text, method)。
    method: text_layer | ocr | empty
    """
    text = _extract_pdf_text_layer(path, max_chars)
    method = "text_layer" if text else "empty"
    min_chars = _ocr_min_text_chars()
    stripped_len = len(re.sub(r"\s+", "", text or ""))

    # 文字层足够：明确不算 OCR（供进度文案）
    if text and stripped_len >= min_chars:
        if on_progress:
            try:
                # page=0 表示文字层路径
                on_progress(0, 0, path.name)
            except Exception:
                pass
        return text, method

    need_ocr = (not text) or stripped_len < min_chars
    if need_ocr and _ocr_enabled():
        logger.info("PDF text sparse (%s chars), trying OCR: %s", stripped_len, path.name)
        ocr_text = ocr_pdf_pages(path, max_chars, on_progress=on_progress)
        if ocr_text and len(re.sub(r"\s+", "", ocr_text)) > stripped_len:
            return ocr_text, "ocr"
        if ocr_text and not text:
            return ocr_text, "ocr"

    if text:
        return text, method

    return (
        "（PDF 未能提取到有效文字。"
        "已尝试 OCR，但仍无结果。请确认 tesseract 与中文语言包已安装，"
        "或上传可复制文字的 PDF / DOCX / TXT。）"
    ), "empty"


def extract_docx(path: Path, max_chars: int = 14000) -> str:
    try:
        import docx  # python-docx
    except Exception as e:
        raise RuntimeError(f"python-docx 不可用: {e}") from e

    document = docx.Document(str(path))
    parts = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _clean("\n".join(parts))
    if not text:
        return "（DOCX 未提取到有效文字）"
    return _truncate(text, max_chars)


def extract_pptx(path: Path, max_chars: int = 14000) -> str:
    try:
        from pptx import Presentation
    except Exception as e:
        raise RuntimeError(f"python-pptx 不可用: {e}") from e

    prs = Presentation(str(path))
    parts: list[str] = []
    total = 0
    for idx, slide in enumerate(prs.slides, 1):
        slide_bits: list[str] = []
        for shape in slide.shapes:
            try:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    t = shape.text_frame.text
                    if t and t.strip():
                        slide_bits.append(t.strip())
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                        if cells:
                            slide_bits.append(" | ".join(cells))
            except Exception:
                continue
        if slide_bits:
            chunk = f"\n--- 第 {idx} 页 ---\n" + "\n".join(slide_bits)
            parts.append(chunk)
            total += len(chunk)
            if total >= max_chars:
                break
    text = _clean("\n".join(parts))
    if not text:
        return "（PPTX 未提取到有效文字）"
    return _truncate(text, max_chars)


def extract_path(
    path: str | Path,
    ext: str | None = None,
    max_chars: int = 14000,
    on_progress=None,
) -> dict:
    p = Path(path)
    if not p.exists() or not p.is_file():
        return {
            "ok": False,
            "error": "file_not_found",
            "text": "",
            "path": str(p),
        }

    suffix = (ext or p.suffix.lstrip(".")).lower()
    method = "native"
    try:
        if suffix in TEXT_EXTS:
            text = extract_text_file(p, max_chars)
            method = "text"
        elif suffix == "pdf":
            text, method = extract_pdf(p, max_chars, on_progress=on_progress)
        elif suffix in {"doc", "docx"}:
            if suffix == "doc":
                return {
                    "ok": False,
                    "error": "unsupported_legacy_doc",
                    "text": "（.doc 旧格式暂不支持，请转换为 .docx 或 .pdf）",
                    "path": str(p),
                }
            text = extract_docx(p, max_chars)
            method = "docx"
        elif suffix == "pptx":
            text = extract_pptx(p, max_chars)
            method = "pptx"
        elif suffix == "ppt":
            return {
                "ok": False,
                "error": "unsupported_legacy_ppt",
                "text": "（.ppt 旧格式暂不支持，请转换为 .pptx）",
                "path": str(p),
            }
        else:
            return {
                "ok": False,
                "error": "unsupported_ext",
                "text": f"（暂不支持解析 .{suffix} 全文，仅提供文件名上下文）",
                "path": str(p),
            }

        ok = bool(text) and not text.startswith("（PDF 未能提取") and not text.startswith("（DOCX 未") and not text.startswith("（PPTX 未")
        if method == "empty":
            ok = False
        return {
            "ok": ok if method != "empty" else False,
            "error": None if method != "empty" else "no_text",
            "text": text,
            "path": str(p),
            "ext": suffix,
            "chars": len(text or ""),
            "method": method,
            "ocr": method == "ocr",
        }
    except Exception as e:
        logger.exception("extract failed: %s", p)
        return {
            "ok": False,
            "error": str(e),
            "text": f"（解析失败：{e}）",
            "path": str(p),
        }
