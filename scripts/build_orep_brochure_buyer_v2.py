from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_orep_brochure import (
    BLUE,
    CYAN,
    GREEN,
    INK,
    LIGHT_BLUE,
    MUTED,
    NAVY,
    OUT_DIR,
    add_bullets,
    add_callout,
    add_heading,
    add_header_footer,
    add_p,
    add_section_divider,
    add_table,
    set_page,
    set_run,
)


OUT_DOCX = OUT_DIR / "OREP在线路演评审平台_产品宣传手册_买方叙事版.docx"


def add_cover(doc: Document) -> None:
    add_header_footer(doc.sections[0], "OREP Product Brochure")
    add_p(doc, "OREP 产品宣传手册", size=10, color=CYAN, bold=True, before=18, after=20)
    p = add_p(doc, "在线路演评审平台", size=31, color=NAVY, bold=True, before=0, after=8)
    p.paragraph_format.line_spacing = 1.05
    add_p(
        doc,
        "面向高校、赛事与创新创业项目评审的一体化数字评审与成长平台",
        size=15,
        color=INK,
        before=0,
        after=18,
    )
    add_callout(
        doc,
        "一句话定位",
        "OREP 帮助组织方把路演评审从一次活动，升级为可组织、可追溯、可复盘、可持续运营的数字化体系。",
        fill=LIGHT_BLUE,
    )
    add_table(
        doc,
        ["谁会关心", "他们最想解决什么"],
        [
            ["赛事/学校管理者", "少靠人工串联，提升组织效率，沉淀过程数据。"],
            ["评委/指导老师", "更清楚地评分、反馈、追踪问题和复盘依据。"],
            ["参赛团队", "不仅拿到分数，还知道 PPT、表达和现场呈现如何改。"],
            ["信息化/IT 团队", "可私有化部署、可治理数据、可长期运营。"],
        ],
        [2.0, 4.35],
        first_col=True,
    )
    add_p(doc, "正式宣传版 | 2026-06-10", size=9, color=MUTED, before=24, after=0)
    doc.add_page_break()


def page_why(doc: Document) -> None:
    add_section_divider(
        doc,
        "01",
        "为什么需要 OREP",
        "路演评审的难点，不只是把会议搬到线上，而是让评审过程真正可管理、可解释、可复用。",
    )
    add_table(
        doc,
        ["常见问题", "对组织方的影响", "OREP 的改变"],
        [
            ["组织分散", "会议、材料、评分、反馈和报告分散在多个工具，活动组织依赖人工协调。", "用统一平台承载赛前准备、赛中评审、赛后复盘和运营管理。"],
            ["评分难追溯", "分数之外缺少扣分依据、问题记录和可归档材料，复核成本高。", "形成评分记录、问题闭环、录制留痕和报告输出。"],
            ["反馈不成体系", "学生只知道结果，不知道表达、材料、逻辑和现场呈现具体怎么改。", "用 AI 画像、表达诊断和训练建议把反馈变成行动。"],
            ["经验难复用", "每次比赛重新组织，模板、资源、历史项目和数据难沉淀。", "建立资源库、历史任务、数据监控和后台治理能力。"],
        ],
        [1.35, 2.5, 2.5],
        first_col=True,
    )
    add_callout(doc, "读者应当记住", "OREP 的核心不是“多一个路演工具”，而是让组织方拥有一套数字化评审基础设施。")
    doc.add_page_break()


def page_what(doc: Document) -> None:
    add_section_divider(
        doc,
        "02",
        "OREP 是什么",
        "一套围绕路演评审全流程设计的平台，把准备、评审、复盘和运营放进同一个闭环。",
    )
    add_table(
        doc,
        ["阶段", "用户动作", "平台交付"],
        [
            ["赛前准备", "学生准备项目材料，教师/组织方提供模板和要求。", "PPT 生成、素材证据、讲稿训练、资源中心。"],
            ["赛中评审", "团队在线路演，评委参与会议并完成评分。", "在线会议、专家评分、聊天互动、录制留痕。"],
            ["赛后复盘", "组织方汇总结果，学生查看反馈并进入下一轮优化。", "评分报告、问题闭环、AI 能力画像、训练建议。"],
            ["持续运营", "学校/赛事/园区沉淀资源、数据、历史项目和管理经验。", "后台治理、数据监控、资源管理、私有化部署。"],
        ],
        [1.2, 2.6, 2.55],
        first_col=True,
    )
    add_heading(doc, "产品不是功能堆叠，而是一个闭环", 2)
    add_bullets(
        doc,
        [
            "对组织方：从报名、准备、评审、汇总到归档，减少跨工具协调。",
            "对评委：评分、评论、问题和结果统一沉淀，提升评审一致性。",
            "对学生：从“看到分数”升级为“看到证据、问题和下一步训练动作”。",
        ]
    )
    doc.add_page_break()


def page_journey(doc: Document) -> None:
    add_section_divider(
        doc,
        "03",
        "一场路演如何运转",
        "用真实工作流解释产品，而不是让读者在功能清单里寻找答案。",
    )
    add_table(
        doc,
        ["角色", "关键动作", "获得的价值"],
        [
            ["组织方", "创建会议、配置资源、管理用户、查看评分与问题。", "活动更易组织，过程更易归档，数据可持续沉淀。"],
            ["参赛团队", "准备 PPT、绑定素材、参加会议、查看 AI 反馈。", "知道材料、表达和演示的具体改进方向。"],
            ["评委", "进入会议、观看路演、按评分项提交分数和评论。", "评分过程更统一，反馈更结构化。"],
            ["管理员", "维护账号角色、资源、会议、PPT 任务和数据监控。", "形成可运营的平台，而不是一次性活动工具。"],
        ],
        [1.15, 2.7, 2.5],
        first_col=True,
    )
    add_callout(doc, "设计原则", "每个功能都服务一个明确场景：组织、评审、复盘、训练或运营。不能解释场景价值的功能，不应成为宣传手册主角。")
    doc.add_page_break()


def page_prepare(doc: Document) -> None:
    add_section_divider(
        doc,
        "04",
        "赛前准备：让材料先具备得分能力",
        "路演质量不是从上台开始，而是从材料结构、证据完整性和讲稿准备开始。",
    )
    add_table(
        doc,
        ["能力", "解决的问题", "用户看到的结果"],
        [
            ["PPT 智能生成", "团队不知道如何把项目材料组织成路演结构。", "生成可预览、可编辑、可下载的路演材料初稿。"],
            ["评分点覆盖", "页面好看但不一定支撑比赛评分标准。", "看到哪些评分点已覆盖，哪些证据需要补强。"],
            ["素材证据管理", "政策、截图、设备照片、数据图等材料难以和页面对应。", "素材与页面、实操步骤、评分点形成绑定。"],
            ["讲稿训练", "PPT 有了，但表达节奏和讲解逻辑仍不稳定。", "形成页面讲稿和训练文本，为彩排做准备。"],
        ],
        [1.45, 2.45, 2.45],
        first_col=True,
    )
    add_callout(doc, "对参赛团队的吸引点", "OREP 不是只帮你做一份 PPT，而是帮你把 PPT 做成能被评委理解、能支撑评分、能用于训练的路演材料。")
    doc.add_page_break()


def page_review(doc: Document) -> None:
    add_section_divider(
        doc,
        "05",
        "赛中评审：让评审过程有序、留痕、可复核",
        "评审现场需要稳定协同，也需要把关键过程转化为后续复盘依据。",
    )
    add_table(
        doc,
        ["能力", "用户体验", "管理价值"],
        [
            ["在线会议", "团队、评委和组织方通过统一入口参与路演。", "减少线下排期和跨工具组织成本。"],
            ["专家评分", "评委按评分项提交分数、评论和判断依据。", "评分口径更集中，结果更容易汇总。"],
            ["互动协同", "会议内支持沟通、状态流转和参会人管理。", "现场管理更可控。"],
            ["录制留痕", "路演过程可保存并用于复盘。", "为争议复核、AI 分析和教学反馈提供素材。"],
        ],
        [1.35, 2.5, 2.5],
        first_col=True,
    )
    add_heading(doc, "不把“会议工具”当终点", 2)
    add_bullets(doc, ["会议只是评审现场，OREP 更重要的是把现场行为转化为评分、证据、报告和成长反馈。"])
    doc.add_page_break()


def page_after(doc: Document) -> None:
    add_section_divider(
        doc,
        "06",
        "赛后复盘：让分数变成改进动作",
        "优秀的评审平台不只给结果，还要帮助学生和指导老师知道下一轮怎么提升。",
    )
    add_table(
        doc,
        ["输出", "说明"],
        [
            ["评分报告", "将评分结果、维度表现和关键问题整理为可归档材料。"],
            ["问题闭环", "把扣分项、评论和待改进事项转化为可跟踪的问题记录。"],
            ["AI 能力画像", "结合转写、表达节奏、现场呈现和评分结果，形成描述性反馈。"],
            ["训练建议", "把表达、材料、演示和证据短板收束为下一轮可执行任务。"],
        ],
        [1.5, 4.85],
        first_col=True,
    )
    add_callout(doc, "AI 的正确讲法", "不要只说“AI 赋能”。手册里应说明 AI 具体产出什么：转写、表达诊断、能力画像、证据链和训练建议。")
    doc.add_page_break()


def page_operation(doc: Document) -> None:
    add_section_divider(
        doc,
        "07",
        "持续运营：让一次比赛沉淀为长期资产",
        "当平台持续承载比赛、课程和训练营，组织方才能积累资源、经验和数据。",
    )
    add_table(
        doc,
        ["运营对象", "平台支持", "沉淀资产"],
        [
            ["用户与角色", "维护管理员、教师、学生、评委、专家等角色。", "清晰的访问身份和组织权限。"],
            ["资源与模板", "统一管理 PPT 模板、评分表、交付文档和素材。", "可复用的赛事与课程资源库。"],
            ["会议与任务", "管理会议、评分、PPT 任务、历史记录和问题。", "完整的活动过程档案。"],
            ["数据监控", "查看在线状态、行为日志和平台使用情况。", "运营决策和服务改进依据。"],
        ],
        [1.4, 2.45, 2.5],
        first_col=True,
    )
    doc.add_page_break()


def page_scenarios(doc: Document) -> None:
    add_section_divider(
        doc,
        "08",
        "典型应用场景",
        "同一套评审闭环，可以服务多种路演、答辩和项目评估场景。",
    )
    add_table(
        doc,
        ["场景", "使用方式", "吸引点"],
        [
            ["创新创业大赛", "赛前准备材料，赛中专家评分，赛后输出报告和训练建议。", "提升组织效率和评分规范性。"],
            ["课程答辩/项目验收", "教师创建会议，学生在线演示，平台留存评分与问题。", "让教学评价可追溯、可复盘。"],
            ["园区/孵化器路演", "为项目提供材料打磨、专家评审和复盘服务。", "提升项目展示质量和投融资准备度。"],
            ["校内训练营", "多轮路演训练，每轮保留评分、画像和改进任务。", "把比赛准备变成持续训练机制。"],
            ["远程专家评审", "跨地域专家进入会议并完成标准化评分。", "降低专家组织成本，扩大评审覆盖。"],
        ],
        [1.4, 2.65, 2.3],
        first_col=True,
    )
    doc.add_page_break()


def page_trust(doc: Document) -> None:
    add_section_divider(
        doc,
        "09",
        "为什么可信",
        "技术信息在宣传手册中应转化为采购和部署能理解的信任语言。",
    )
    add_table(
        doc,
        ["可信维度", "买方关心的问题", "OREP 的回答"],
        [
            ["流程可信", "评审过程是否可追溯、可复核？", "会议、评分、问题、报告和历史记录统一沉淀。"],
            ["证据可信", "AI 建议是否有依据？", "反馈来自转写、表达节奏、现场呈现、材料证据和评分结果。"],
            ["数据可信", "校内数据是否可控？", "支持私有化部署和组织内数据治理，便于校内或机构环境落地。"],
            ["交付可信", "能否稳定运维和持续扩展？", "采用模块化服务、统一入口、文件管理、数据存储和后台监控能力。"],
        ],
        [1.35, 2.4, 2.6],
        first_col=True,
    )
    add_callout(
        doc,
        "技术内容边界",
        "宣传手册只需要说明部署形态、数据控制、实时协同和 AI 分析能力；具体框架、接口、端口、数据库表和配置方式应放在技术白皮书或部署文档。",
        fill="EAFBF8",
        color=GREEN,
    )
    doc.add_page_break()


def page_get_started(doc: Document) -> None:
    add_section_divider(
        doc,
        "10",
        "如何落地",
        "从小规模试运行开始，让组织方先验证流程，再扩大到赛事、课程和训练营。",
    )
    add_table(
        doc,
        ["步骤", "建议动作", "阶段成果"],
        [
            ["1. 明确场景", "选择一类先行场景，如创新创业大赛初赛或课程答辩。", "确定评分表、参与角色和流程边界。"],
            ["2. 配置资源", "准备模板、评分项、会议安排、账号角色和材料要求。", "形成可执行的评审空间。"],
            ["3. 试运行", "用一场小规模路演验证会议、评分、报告和复盘。", "发现流程问题并优化组织方式。"],
            ["4. 正式运营", "推广到更多项目、班级、赛事或园区活动。", "沉淀资源、历史数据和持续训练机制。"],
        ],
        [1.05, 3.05, 2.25],
        first_col=True,
    )
    add_callout(
        doc,
        "最终价值",
        "OREP 让路演评审从“组织一场活动”，升级为“建设一套可持续运营的数字化评审与成长体系”。",
    )


def build_doc() -> Document:
    doc = Document()
    set_page(doc)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["List Bullet"].font.name = "Microsoft YaHei"
    styles["List Bullet"].font.size = Pt(10.5)
    add_cover(doc)
    page_why(doc)
    page_what(doc)
    page_journey(doc)
    page_prepare(doc)
    page_review(doc)
    page_after(doc)
    page_operation(doc)
    page_scenarios(doc)
    page_trust(doc)
    page_get_started(doc)
    doc.core_properties.title = "OREP 在线路演评审平台产品宣传手册 - 买方叙事版"
    doc.core_properties.subject = "基于软件产品宣传手册 benchmark 重写"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "技术细节已按宣传手册规则收敛为可信交付语言"
    doc.core_properties.created = datetime.now()
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
