from __future__ import annotations

from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_orep_brochure import (
    BLUE,
    CYAN,
    GOLD,
    GREEN,
    INK,
    LIGHT_BLUE,
    MUTED,
    NAVY,
    OUT_DIR,
    add_callout,
    add_header_footer,
    add_p,
    add_section_divider,
    add_table,
    set_run,
)
from build_orep_brochure_word_v3 import (
    P,
    add_card_grid,
    add_full_page_image,
    add_image,
    configure_a4_section,
    draw_text,
    rounded,
)


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = OUT_DIR / "smart_brain_assets"
OUT_DOCX = OUT_DIR / "智慧大脑_产品宣传手册_融合规划版.docx"


def make_cover(path: Path) -> None:
    w, h = 1600, 2264
    img = Image.new("RGB", (w, h), P["navy"])
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, w, h), fill=P["navy"])
    d.rectangle((0, 0, 170, h), fill="#08265F")
    d.rectangle((170, 0, 214, h), fill=P["bright"])
    rounded(d, (275, 200, 690, 262), 18, "#092B68", "#1B7CFF", 2)
    draw_text(d, (306, 218), "SMART BRAIN", 30, P["cyan"], True)
    draw_text(d, (275, 365), "智慧大脑", 106, "#FFFFFF", True)
    draw_text(d, (282, 525), "创新创业路演评审与成长训练平台", 44, "#DCEBFF", True)
    draw_text(
        d,
        (282, 682),
        "让每一次路演都成为下一次成长的证据",
        46,
        "#FFFFFF",
        True,
    )
    draw_text(
        d,
        (282, 785),
        "融合路演评审、AI 复盘、21 天训练、能力档案与组织运营，帮助学校、赛事和园区建立持续成长机制。",
        34,
        "#BFD8FF",
        False,
        max_chars=29,
        line_gap=16,
    )

    rounded(d, (275, 1040, 1390, 1575), 34, "#F4F7FB", "#2F66C9", 3)
    cards = [
        ("评审", "路演、评分、留痕", P["blue"]),
        ("复盘", "报告、问题、AI 画像", P["green"]),
        ("训练", "21 天任务、课程、考试", P["amber"]),
        ("成长", "能力档案、证书、趋势", P["cyan"]),
    ]
    for i, (title, body, color) in enumerate(cards):
        x = 340 + (i % 2) * 510
        y = 1110 + (i // 2) * 200
        rounded(d, (x, y, x + 405, y + 135), 26, "#FFFFFF", "#D7E2F0", 3)
        d.rectangle((x, y, x + 14, y + 135), fill=color)
        draw_text(d, (x + 46, y + 30), title, 42, color, True)
        draw_text(d, (x + 46, y + 90), body, 27, P["muted"])

    rounded(d, (275, 1715, 1390, 1858), 28, "#092B68", "#244F96", 2)
    draw_text(d, (320, 1752), "不是多一个工具，而是一套把评审结果转化为训练行动的成长系统。", 32, "#DCEBFF", False, max_chars=33)
    draw_text(d, (275, 2115), "融合规划版  |  2026", 28, "#8EABD8", True)
    img.save(path)


def make_growth_loop(path: Path) -> None:
    w, h = 1500, 860
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    rounded(d, (45, 45, w - 45, h - 45), 32, "#F8FBFF", "#D7E2F0", 3)
    d.ellipse((620, 305, 880, 565), fill=P["pale"], outline="#C5D7EE", width=4)
    d.ellipse((655, 340, 845, 530), fill=P["navy"])
    draw_text(d, (690, 388), "智慧\n大脑", 42, "#FFFFFF", True, line_gap=10)
    nodes = [
        ("赛前准备", "材料、讲稿、评分点", 115, 120, P["blue"]),
        ("在线路演", "会议、演示、互动", 570, 105, P["green"]),
        ("专家评审", "分数、评论、依据", 1055, 120, P["blue"]),
        ("AI 复盘", "报告、问题、画像", 1055, 585, P["amber"]),
        ("21 天训练", "课程、任务、考试", 570, 610, P["green"]),
        ("能力档案", "雷达、趋势、证书", 115, 585, P["blue"]),
    ]
    for _, _, x, y, _ in nodes:
        d.line((x + 160, y + 61, 750, 435), fill="#BFD0E7", width=4)
    for title, body, x, y, color in nodes:
        rounded(d, (x, y, x + 320, y + 122), 24, "#FFFFFF", "#D7E2F0", 3)
        d.rectangle((x, y, x + 12, y + 122), fill=color)
        draw_text(d, (x + 40, y + 28), title, 32, color, True)
        draw_text(d, (x + 40, y + 75), body, 24, P["muted"])
    img.save(path)


def make_21day_path(path: Path) -> None:
    w, h = 1500, 860
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    rounded(d, (45, 45, w - 45, h - 45), 32, "#F8FBFF", "#D7E2F0", 3)
    draw_text(d, (95, 88), "21 天训练把评审反馈变成每日行动", 42, P["navy"], True)
    phases = [
        ("第 1-3 天", "理解问题", "查看复盘报告、明确短板和训练目标", P["blue"]),
        ("第 4-10 天", "补齐基础", "课程学习、资源阅读、每日任务提交", P["green"]),
        ("第 11-17 天", "强化表达", "讲稿训练、打字速度、演示节奏优化", P["amber"]),
        ("第 18-21 天", "验证成果", "在线考试、团队评分、能力档案更新", P["cyan"]),
    ]
    for i, (day, title, body, color) in enumerate(phases):
        x = 110 + i * 340
        rounded(d, (x, 230, x + 280, 590), 28, "#FFFFFF", "#D7E2F0", 3)
        d.rectangle((x, 230, x + 280, 242), fill=color)
        draw_text(d, (x + 30, 285), day, 28, color, True)
        draw_text(d, (x + 30, 345), title, 40, P["navy"], True)
        draw_text(d, (x + 30, 435), body, 25, P["muted"], False, max_chars=10, line_gap=10)
        if i < 3:
            d.line((x + 290, 410, x + 330, 410), fill=P["bright"], width=5)
            d.polygon([(x + 330, 398), (x + 352, 410), (x + 330, 422)], fill=P["bright"])
    rounded(d, (110, 675, 1390, 770), 24, P["green_pale"], "#B8E8DE", 3)
    draw_text(d, (150, 703), "训练不是附属功能，而是让评审结果真正产生改变的关键环节。", 30, P["green"], True)
    img.save(path)


def make_capability_map(path: Path) -> None:
    w, h = 1500, 860
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    rounded(d, (45, 45, w - 45, h - 45), 32, "#F8FBFF", "#D7E2F0", 3)
    groups = [
        ("组织运营", ["用户/部门", "权限/日志", "课程/资源"], P["blue"]),
        ("评审现场", ["在线会议", "专家评分", "团队路演"], P["green"]),
        ("训练学习", ["每日任务", "课程学习", "移动学习"], P["amber"]),
        ("测评认证", ["在线考试", "AI 分析", "证书档案"], P["cyan"]),
    ]
    for i, (title, items, color) in enumerate(groups):
        x = 105 + (i % 2) * 655
        y = 105 + (i // 2) * 330
        rounded(d, (x, y, x + 560, y + 245), 28, "#FFFFFF", "#D7E2F0", 3)
        d.rectangle((x, y, x + 14, y + 245), fill=color)
        draw_text(d, (x + 44, y + 34), title, 42, color, True)
        for j, item in enumerate(items):
            rounded(d, (x + 48 + j * 158, y + 132, x + 178 + j * 158, y + 184), 14, P["pale"], "#D7E2F0", 2)
            draw_text(d, (x + 72 + j * 158, y + 146), item, 22, P["ink"], True)
    img.save(path)


def make_closing(path: Path) -> None:
    w, h = 1600, 2264
    img = Image.new("RGB", (w, h), P["navy"])
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, w, h), fill=P["navy"])
    d.rectangle((0, h - 400, w, h), fill="#071735")
    draw_text(d, (190, 290), "从一次路演", 86, "#FFFFFF", True)
    draw_text(d, (190, 430), "到一套持续成长机制", 82, "#FFFFFF", True)
    draw_text(
        d,
        (195, 680),
        "智慧大脑的目标，不是让组织方多管理一个系统，而是让每一次评审都留下可训练、可复盘、可认证的成长资产。",
        38,
        "#DCEBFF",
        False,
        max_chars=28,
        line_gap=16,
    )
    for i, (title, body) in enumerate([
        ("可评审", "路演过程有标准"),
        ("可复盘", "反馈结果有证据"),
        ("可训练", "问题变成每日行动"),
        ("可认证", "成长沉淀为档案"),
    ]):
        x = 190 + (i % 2) * 590
        y = 995 + (i // 2) * 235
        rounded(d, (x, y, x + 500, y + 145), 28, "#092B68", "#244F96", 3)
        draw_text(d, (x + 38, y + 30), title, 40, P["cyan"], True)
        draw_text(d, (x + 38, y + 88), body, 28, "#DCEBFF")
    rounded(d, (190, 1610, 1410, 1800), 34, "#FFFFFF", "#C7D7EF", 3)
    draw_text(d, (235, 1652), "建议下一步", 38, P["blue"], True)
    draw_text(d, (235, 1715), "先以一场真实路演评审试运行，再接入 21 天训练模块，验证从评审到成长的完整闭环。", 30, P["ink"], False, max_chars=34)
    draw_text(d, (190, 2052), "适用于：高校创新创业学院、赛事承办方、园区/孵化器、课程答辩与项目训练营", 30, "#AFC8EC", False)
    img.save(path)


def create_assets() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_cover(ASSET_DIR / "cover.png")
    make_growth_loop(ASSET_DIR / "growth_loop.png")
    make_21day_path(ASSET_DIR / "day21_path.png")
    make_capability_map(ASSET_DIR / "capability_map.png")
    make_closing(ASSET_DIR / "closing.png")


def new_doc() -> Document:
    doc = Document()
    configure_a4_section(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["List Bullet"].font.name = "Microsoft YaHei"
    styles["List Bullet"].font.size = Pt(10.5)
    add_header_footer(doc.sections[0], "智慧大脑 Product Brochure")
    return doc


def page_why(doc: Document) -> None:
    add_section_divider(doc, "01", "为什么不能只做一次路演", "评审结束不是成长结束。真正的问题，是结果如何被理解、训练和沉淀。")
    add_card_grid(doc, [
        ("评审结束即结束", "传统路演常停留在现场表现和最终分数，学生知道结果，却不知道下一轮如何改。", BLUE),
        ("反馈难以落地", "评委意见、材料问题、表达短板和现场表现分散存在，难以转成持续训练任务。", GREEN),
        ("成长不可见", "学校很难持续看到学生能力变化，也难以沉淀跨赛事、跨课程的成长档案。", GOLD),
        ("组织经验难复用", "每次赛事重新组织，课程、题库、资源、证书和历史数据没有形成运营资产。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "手册核心观点", "智慧大脑不是把功能堆得更多，而是把“评审结果”变成“训练行动”，再沉淀为“成长证据”。")
    doc.add_page_break()


def page_position(doc: Document) -> None:
    add_section_divider(doc, "02", "智慧大脑是什么", "它不是单一评审工具，而是围绕创新创业训练建立的评审、复盘、训练、测评与档案闭环。")
    add_image(doc, ASSET_DIR / "growth_loop.png", width=7.0, before=4, after=6)
    add_callout(doc, "产品定位", "面向高校、赛事、园区与创新创业训练场景，智慧大脑帮助组织方把一次性路演活动升级为可评审、可复盘、可训练、可认证、可持续运营的成长体系。")
    doc.add_page_break()


def page_flow(doc: Document) -> None:
    add_section_divider(doc, "03", "一条完整成长路径", "用户不需要记住所有功能，只需要看懂一个变化：学生从被评分，走向被训练、被看见。")
    add_table(
        doc,
        ["阶段", "用户动作", "平台交付"],
        [
            ["赛前准备", "团队准备项目材料、讲稿和评分点证据。", "PPT 材料、资源模板、讲稿训练、评分点覆盖。"],
            ["在线路演", "团队演示，评委进入会议并提交评分。", "会议协同、专家评分、团队路演评分、录制留痕。"],
            ["赛后复盘", "组织方汇总结果，学生查看问题和建议。", "评分报告、AI 画像、问题闭环、训练建议。"],
            ["21 天训练", "学生按计划学习课程、完成每日任务和测评。", "课程学习、每日任务、在线考试、打字速度记录。"],
            ["能力沉淀", "教师和学生查看能力变化和认证成果。", "能力雷达、评分趋势、证书、成长档案。"],
        ],
        [1.2, 2.6, 2.55],
        first_col=True,
    )
    doc.add_page_break()


def page_before(doc: Document) -> None:
    add_section_divider(doc, "04", "赛前：让项目先具备表达能力", "好的路演不是上台才开始，而是从材料结构、证据完整性和表达训练开始。")
    add_card_grid(doc, [
        ("材料组织", "围绕项目亮点、应用价值、评分点和证据材料组织路演内容。", BLUE),
        ("讲稿训练", "把页面内容转成可演练、可迭代的讲稿，降低临场表达不稳定。", GREEN),
        ("资源模板", "沉淀课程、课件、视频、图片和项目素材，减少每次重新准备。", GOLD),
        ("评分预期", "提前明确评审维度，让团队知道评委真正看什么。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "给学生的价值", "不是只帮学生做一份材料，而是让学生知道如何把项目讲清楚、讲可信、讲到评审点上。")
    doc.add_page_break()


def page_review(doc: Document) -> None:
    add_section_divider(doc, "05", "赛中：让评审过程可组织、可追溯", "现场评审既要顺畅协同，也要把关键过程转化为后续复盘依据。")
    add_card_grid(doc, [
        ("在线会议", "团队、评委、组织方通过统一入口参与路演，降低跨地域组织成本。", BLUE),
        ("专家评分", "按评分标准记录分数、评论和判断依据，便于汇总与复核。", GREEN),
        ("团队路演评分", "支持多维度评分与历史记录，持续观察团队表现变化。", GOLD),
        ("录制留痕", "保留现场过程，为争议复核、教学反馈和 AI 分析提供素材基础。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_after(doc: Document) -> None:
    add_section_divider(doc, "06", "赛后：让分数变成训练任务", "真正有价值的复盘，不是告诉学生得了多少分，而是告诉他下一轮如何变好。")
    add_card_grid(doc, [
        ("评分报告", "将总分、维度表现和关键问题整理为可归档材料。", BLUE),
        ("问题闭环", "把扣分点、评委评论和待改进事项转成可跟踪问题。", GREEN),
        ("AI 能力画像", "结合转写、表达节奏、现场呈现和评分结果形成描述性反馈。", GOLD),
        ("训练建议", "把表达、材料、演示和证据短板收束为 21 天训练任务。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "表达边界", "AI 不应被写成空泛口号。宣传册里只讲清楚它帮助生成什么：复盘报告、能力画像、问题定位和训练建议。")
    doc.add_page_break()


def page_training(doc: Document) -> None:
    add_section_divider(doc, "07", "21 天训练：让反馈真正落地", "21 天训练模块是融合规划中的成长训练能力，用来承接路演复盘后的持续提升。")
    add_image(doc, ASSET_DIR / "day21_path.png", width=7.0, before=4, after=6)
    add_callout(doc, "融合口径", "21 天训练当前作为智慧大脑的融合规划模块表达：它可以承接复盘建议，组织课程、每日任务、考试测评和学习记录，但不应写成已经完全融合上线。")
    doc.add_page_break()


def page_learning(doc: Document) -> None:
    add_section_divider(doc, "08", "学习与任务：把成长拆成每天能完成的动作", "训练不是一句建议，而是课程、任务、记录和进度共同构成的学习路径。")
    add_card_grid(doc, [
        ("课程学习", "课程、章节、课时、视频、课件附件构成结构化训练内容。", BLUE),
        ("每日任务", "支持每日文字任务提交、训练记录保存和历史查看。", GREEN),
        ("打字速度", "记录打字速度、时间和字数，形成日常训练数据。", GOLD),
        ("移动学习", "PC 与 H5 移动端数据同步，便于学生随时继续学习。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_exam(doc: Document) -> None:
    add_section_divider(doc, "09", "考试与测评：让训练结果可验证", "训练后的改变需要被测量，测评结果也要继续反哺下一轮训练。")
    add_card_grid(doc, [
        ("在线考试", "支持考试创建、发布、分配、作答、提交和成绩查看。", BLUE),
        ("题库刷题", "题目支持单选、多选、判断、填空和编程题，帮助形成练习体系。", GREEN),
        ("自动与人工批改", "客观题自动判分，主观题和编程题可人工批改与评语反馈。", GOLD),
        ("AI 成绩分析", "基于作答数据生成成绩分析和个性化学习建议，依赖系统 AI 配置启用。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_profile(doc: Document) -> None:
    add_section_divider(doc, "10", "能力档案：让成长被看见", "智慧大脑要卖的不是分数，而是让学生、教师和组织方共同看见成长轨迹。")
    add_card_grid(doc, [
        ("六维雷达", "展示解决问题、代码、沟通、团队协作、演讲、创意等能力分布。", BLUE),
        ("评分趋势", "展示团队路演总分和最近评分历史，观察能力变化。", GREEN),
        ("评级标签", "根据平均能力分数生成评级标签，形成直观能力标识。", GOLD),
        ("证书沉淀", "支持证书模板、颁发、查看与下载，让学习成果可认证。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "品牌层表达", "客户会忘记功能名，但会记住一件事：智慧大脑让学生看见一个自己还没看见的可能性。")
    doc.add_page_break()


def page_admin(doc: Document) -> None:
    add_section_divider(doc, "11", "组织运营后台：让能力建设可持续", "当课程、资源、学员、部门、评分和考试统一沉淀，学校才能把训练变成长期资产。")
    add_image(doc, ASSET_DIR / "capability_map.png", width=7.0, before=4, after=6)
    add_table(
        doc,
        ["对象", "管理内容", "组织价值"],
        [
            ["用户与部门", "学员、教师、管理员、部门树、批量导入。", "适配学校、企业、园区组织结构。"],
            ["课程与资源", "课程、章节、课时、视频、课件、分类管理。", "形成可复用训练资源库。"],
            ["权限与审计", "角色权限、数据权限、管理日志、账号集成。", "保障数据边界与运营安全。"],
        ],
        [1.25, 2.65, 2.45],
        first_col=True,
    )
    doc.add_page_break()


def page_scenarios(doc: Document) -> None:
    add_section_divider(doc, "12", "典型应用场景", "同一套成长闭环，可以服务赛事、课程、园区和训练营等多种组织场景。")
    add_table(
        doc,
        ["场景", "使用方式", "用户记住什么"],
        [
            ["创新创业大赛", "赛前准备、在线路演、专家评审、赛后训练。", "比赛不是终点，而是下一轮成长的起点。"],
            ["课程答辩/项目验收", "教师创建任务，学生演示，系统沉淀评分与问题。", "教学评价可追溯，改进动作可跟踪。"],
            ["校内训练营", "多轮路演结合 21 天训练和考试测评。", "训练过程被看见，能力变化被证明。"],
            ["园区/孵化器", "为项目提供材料打磨、专家评审和成长档案。", "项目服务从一次辅导升级为持续陪跑。"],
            ["企业技能提升", "课程、考试、能力档案和证书联动。", "培训不只看完成率，也看能力提升。"],
        ],
        [1.4, 2.55, 2.4],
        first_col=True,
    )
    doc.add_page_break()


def page_trust(doc: Document) -> None:
    add_section_divider(doc, "13", "为什么可信", "技术参数在宣传册中应转译成采购、部署和运营能理解的信任语言。")
    add_card_grid(doc, [
        ("数据可信", "学习记录、评分记录、考试成绩、资源和证书统一沉淀，可追溯。", BLUE),
        ("内容可信", "支持防刷课、跑马灯水印、防录屏提醒等学习质量保障机制。", GREEN),
        ("组织可信", "支持角色权限、数据权限、管理日志和组织账号集成。", GOLD),
        ("交付可信", "支持多端访问、对象存储、私有化部署和系统配置能力。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "技术内容边界", "正文不展开 React、Spring Boot、MySQL、MinIO、Docker 等技术栈；这些应放在技术白皮书或部署文档。宣传册只保留买方能理解的能力类别。")
    doc.add_page_break()


def page_close(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_a4_section(section)
    section.different_first_page_header_footer = True
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""
    add_full_page_image(doc, ASSET_DIR / "closing.png", width=7.03)


def build_doc() -> Document:
    create_assets()
    doc = new_doc()
    add_full_page_image(doc, ASSET_DIR / "cover.png", width=7.03)
    doc.add_page_break()
    page_why(doc)
    page_position(doc)
    page_flow(doc)
    page_before(doc)
    page_review(doc)
    page_after(doc)
    page_training(doc)
    page_learning(doc)
    page_exam(doc)
    page_profile(doc)
    page_admin(doc)
    page_scenarios(doc)
    page_trust(doc)
    page_close(doc)
    doc.core_properties.title = "智慧大脑产品宣传手册 - 融合规划版"
    doc.core_properties.subject = "创新创业路演评审与成长训练平台"
    doc.core_properties.author = "Codex"
    doc.core_properties.created = datetime.now()
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
