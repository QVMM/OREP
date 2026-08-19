from __future__ import annotations

from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_orep_brochure import (
    BLUE,
    GOLD,
    GREEN,
    INK,
    MUTED,
    NAVY,
    OUT_DIR,
    add_callout,
    add_header_footer,
    add_p,
    add_section_divider,
    add_table,
    set_run,
)
from build_orep_brochure_word_v3 import (
    P,
    add_card_grid,
    add_full_page_image,
    add_image,
    configure_a4_section,
    draw_text,
    rounded,
)


ASSET_DIR = OUT_DIR / "smart_brain_image2_assets"
OUT_DOCX = OUT_DIR / "智慧大脑_产品宣传手册_视觉增强版V3.docx"


def crop_fill(img: Image.Image, size: tuple[int, int]) -> Image.Image:
    tw, th = size
    sw, sh = img.size
    scale = max(tw / sw, th / sh)
    resized = img.resize((int(sw * scale), int(sh * scale)), Image.Resampling.LANCZOS)
    x = (resized.width - tw) // 2
    y = (resized.height - th) // 2
    return resized.crop((x, y, x + tw, y + th))


def make_cover() -> None:
    w, h = 1600, 2264
    src = Image.open(ASSET_DIR / "ai_hero_brain.png").convert("RGB")
    top = crop_fill(src, (w, 1260))
    top = ImageEnhance.Brightness(top).enhance(0.72)
    img = Image.new("RGB", (w, h), P["navy"])
    img.paste(top, (0, 0))
    d = ImageDraw.Draw(img, "RGBA")
    d.rectangle((0, 0, w, 1260), fill=(6, 27, 68, 72))
    d.rectangle((0, 1130, w, h), fill=(6, 27, 68, 255))
    d.rectangle((0, 0, 160, h), fill=(8, 38, 95, 255))
    d.rectangle((160, 0, 206, h), fill=(27, 124, 255, 255))
    rounded(d, (270, 150, 640, 215), 20, "#092B68", "#1B7CFF", 2)
    draw_text(d, (302, 168), "SMART BRAIN", 30, P["cyan"], True)
    draw_text(d, (270, 330), "智慧大脑", 112, "#FFFFFF", True)
    draw_text(d, (275, 505), "创新创业路演评审与成长训练平台", 45, "#DCEBFF", True)
    draw_text(d, (275, 655), "让每一次路演都成为下一次成长的证据", 46, "#FFFFFF", True)
    draw_text(
        d,
        (275, 790),
        "把路演评审、AI 复盘、21 天训练、能力档案与组织运营连接起来，帮助学校、赛事和园区建立持续成长机制。",
        34,
        "#BFD8FF",
        False,
        max_chars=29,
        line_gap=16,
    )
    cards = [
        ("可评审", "路演过程有标准"),
        ("可复盘", "反馈结果有证据"),
        ("可训练", "问题转成每日行动"),
        ("可认证", "成长沉淀为档案"),
    ]
    for i, (title, body) in enumerate(cards):
        x = 270 + (i % 2) * 520
        y = 1430 + (i // 2) * 215
        rounded(d, (x, y, x + 440, y + 145), 28, "#0C347A", "#2B65C9", 3)
        draw_text(d, (x + 38, y + 31), title, 40, P["cyan"], True)
        draw_text(d, (x + 38, y + 91), body, 27, "#DCEBFF")
    draw_text(d, (270, 2118), "视觉增强版 V3  |  融合规划口径  |  2026", 27, "#8EABD8", True)
    img.save(ASSET_DIR / "cover_v2.png")


def make_close() -> None:
    w, h = 1600, 2264
    src = Image.open(ASSET_DIR / "ai_dashboard.png").convert("RGB")
    top = crop_fill(src, (w, 1100))
    top = ImageEnhance.Brightness(top).enhance(0.62)
    img = Image.new("RGB", (w, h), P["navy"])
    img.paste(top, (0, 0))
    d = ImageDraw.Draw(img, "RGBA")
    d.rectangle((0, 0, w, 1100), fill=(6, 27, 68, 105))
    d.rectangle((0, 960, w, h), fill=(6, 27, 68, 255))
    draw_text(d, (190, 310), "从一次路演", 82, "#FFFFFF", True)
    draw_text(d, (190, 450), "到一套持续成长机制", 82, "#FFFFFF", True)
    draw_text(
        d,
        (195, 1160),
        "智慧大脑不是多一个系统，而是把每一次评审留下来的问题、证据和反馈，转化为可训练、可复盘、可认证的成长资产。",
        38,
        "#DCEBFF",
        False,
        max_chars=29,
        line_gap=16,
    )
    rounded(d, (195, 1545, 1405, 1740), 34, "#FFFFFF", "#C7D7EF", 3)
    draw_text(d, (240, 1587), "建议下一步", 38, P["blue"], True)
    draw_text(d, (240, 1650), "先选择一场真实路演评审试运行，再接入 21 天训练模块，验证从评审到成长的完整闭环。", 30, P["ink"], False, max_chars=34)
    draw_text(d, (195, 2050), "适用于：高校创新创业学院、赛事承办方、园区/孵化器、课程答辩与项目训练营", 30, "#AFC8EC", False)
    img.save(ASSET_DIR / "closing_v2.png")


def add_page_header(doc: Document, idx: str, title: str, subtitle: str) -> None:
    add_section_divider(doc, idx, title, subtitle)


def new_doc() -> Document:
    doc = Document()
    configure_a4_section(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    add_header_footer(doc.sections[0], "智慧大脑 Product Brochure")
    return doc


def page_brand(doc: Document) -> None:
    add_page_header(doc, "01", "客户真正会记住什么", "参数会被忘掉，但身份感、价值观和未来可能性会被记住。")
    add_callout(
        doc,
        "品牌层主张",
        "智慧大脑不卖“更多功能”，而是帮助组织方把学生从“被评分的人”变成“能看见下一次成长路径的人”。",
        fill="EAF3FF",
        color=BLUE,
    )
    add_card_grid(doc, [
        ("不是评审结束", "而是下一轮成长开始。", BLUE),
        ("不是分数归档", "而是问题、证据和训练任务被沉淀。", GREEN),
        ("不是功能堆叠", "而是评审、训练、测评、档案围绕同一个成长目标运转。", GOLD),
        ("不是讲参数", "而是让客户相信：每个项目都值得被看见、被训练、被提升。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_loop(doc: Document) -> None:
    add_page_header(doc, "02", "一张图看懂智慧大脑", "把评审、复盘、训练、测评和档案连成一个成长飞轮。")
    add_image(doc, ASSET_DIR / "ai_growth_loop.png", width=7.0, before=2, after=6)
    add_card_grid(doc, [
        ("赛前准备", "材料、讲稿、评分点、资源模板。", BLUE),
        ("在线路演", "会议协同、专家评分、录制留痕。", GREEN),
        ("AI 复盘", "报告、问题、能力画像、训练建议。", GOLD),
        ("成长沉淀", "21 天训练、考试测评、能力档案、证书。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_flow(doc: Document) -> None:
    add_page_header(doc, "03", "从路演现场到成长档案", "客户不需要记住所有模块，只要看懂这条路径。")
    add_table(
        doc,
        ["阶段", "用户看见什么", "智慧大脑留下什么"],
        [
            ["赛前", "项目材料、讲稿、评分点被提前组织。", "可训练的表达基础。"],
            ["赛中", "路演过程、评分依据、评委意见被记录。", "可追溯的评审证据。"],
            ["赛后", "分数、问题、AI 画像和建议被汇总。", "可执行的训练任务。"],
            ["训练", "21 天课程、每日任务、考试测评持续推进。", "可验证的成长过程。"],
            ["沉淀", "雷达图、趋势、证书、评级标签形成档案。", "可复用的组织资产。"],
        ],
        [1.05, 2.75, 2.55],
        first_col=True,
    )
    add_callout(doc, "融合边界", "21 天训练当前作为融合规划模块表达，不能写成已经完全融合上线；宣传册里用“可接入、可承接、规划模块”保持可信。")
    doc.add_page_break()


def page_review(doc: Document) -> None:
    add_page_header(doc, "04", "评审不是终点", "评委给出的每一个分数和问题，都应该成为下一轮训练的输入。")
    add_card_grid(doc, [
        ("赛前准备", "PPT 材料、素材证据、讲稿训练和评分点覆盖，让项目先具备表达能力。", BLUE),
        ("在线评审", "会议、演示、专家评分、团队路演评分和录制留痕，让过程可组织可复核。", GREEN),
        ("AI 复盘", "将转写、表达、现场呈现和评分结果收束为报告、问题和训练建议。", GOLD),
        ("问题闭环", "把扣分项和评委意见变成任务，而不是停留在一次性反馈。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_training(doc: Document) -> None:
    add_page_header(doc, "05", "21 天训练让反馈落地", "把抽象建议拆成每天可完成、可记录、可验证的训练动作。")
    add_image(doc, ASSET_DIR / "ai_21day_journey.png", width=7.0, before=2, after=6)
    add_card_grid(doc, [
        ("理解问题", "查看复盘报告，明确短板与训练目标。", BLUE),
        ("补齐基础", "课程学习、资源阅读、每日任务提交。", GREEN),
        ("强化表达", "讲稿训练、打字速度、演示节奏优化。", GOLD),
        ("验证成果", "在线考试、团队评分、能力档案更新。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_exam(doc: Document) -> None:
    add_page_header(doc, "06", "训练结果要被验证", "成长不是完成率，而是能力是否真的发生变化。")
    add_card_grid(doc, [
        ("在线考试", "考试创建、发布、分配、作答、提交和成绩查看。", BLUE),
        ("题库练习", "单选、多选、判断、填空、编程题，构成练习体系。", GREEN),
        ("批改反馈", "客观题自动判分，主观题和编程题支持人工批改与评语。", GOLD),
        ("AI 分析", "基于作答数据生成成绩分析和个性化学习建议，依赖系统 AI 配置启用。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "买方语言", "不要说“题库管理模块很完整”，要说“训练结果可以被验证，下一轮训练可以有依据”。")
    doc.add_page_break()


def page_profile(doc: Document) -> None:
    add_page_header(doc, "07", "能力档案让成长被看见", "学生、教师和组织方看到的不再只是一个分数，而是一条成长轨迹。")
    add_image(doc, ASSET_DIR / "ai_dashboard.png", width=7.0, before=2, after=6)
    add_card_grid(doc, [
        ("六维雷达", "解决问题、代码、沟通、团队协作、演讲、创意。", BLUE),
        ("评分趋势", "路演总分与历史评分变化，观察团队表现。", GREEN),
        ("证书认证", "证书模板、颁发、查看与下载，沉淀学习成果。", GOLD),
        ("组织运营", "用户、部门、课程、资源、权限和日志统一管理。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_scenarios(doc: Document) -> None:
    add_page_header(doc, "08", "典型应用场景", "同一套成长闭环，可以服务赛事、课程、园区和训练营。")
    add_table(
        doc,
        ["场景", "使用方式", "用户记住什么"],
        [
            ["创新创业大赛", "赛前准备、在线路演、专家评审、赛后训练。", "比赛不是终点，而是下一轮成长的起点。"],
            ["课程答辩", "教师创建任务，学生演示，系统沉淀评分与问题。", "教学评价可追溯，改进动作可跟踪。"],
            ["校内训练营", "多轮路演结合 21 天训练和考试测评。", "训练过程被看见，能力变化被证明。"],
            ["园区/孵化器", "项目材料打磨、专家评审和成长档案。", "项目服务从一次辅导升级为持续陪跑。"],
        ],
        [1.35, 2.55, 2.45],
        first_col=True,
    )
    doc.add_page_break()


def page_trust(doc: Document) -> None:
    add_page_header(doc, "09", "为什么可信", "技术参数要转译成采购、部署和运营能理解的信任语言。")
    add_card_grid(doc, [
        ("数据可信", "学习记录、评分记录、考试成绩、资源和证书统一沉淀，可追溯。", BLUE),
        ("内容可信", "支持防刷课、跑马灯水印、防录屏提醒等学习质量保障机制。", GREEN),
        ("组织可信", "支持角色权限、数据权限、管理日志和组织账号集成。", GOLD),
        ("交付可信", "支持多端访问、对象存储、私有化部署和系统配置能力。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "技术内容边界", "正文不展开 React、Spring Boot、MySQL、MinIO、Docker 等技术栈；这些放在技术白皮书或部署文档。")
    doc.add_page_break()


def page_landing(doc: Document) -> None:
    add_page_header(doc, "10", "如何落地", "先验证一场真实评审，再把复盘结果接入训练模块。")
    add_table(
        doc,
        ["步骤", "建议动作", "阶段成果"],
        [
            ["1. 选场景", "选择一场创新创业比赛、课程答辩或训练营。", "明确角色、评分表和流程边界。"],
            ["2. 跑评审", "完成材料准备、在线路演、专家评分和复盘报告。", "验证评审组织与反馈质量。"],
            ["3. 接训练", "把复盘建议接入 21 天训练任务、课程和测评。", "验证从反馈到行动的闭环。"],
            ["4. 沉档案", "查看能力雷达、趋势、证书和组织数据。", "形成可复用的成长资产。"],
        ],
        [1.05, 3.0, 2.3],
        first_col=True,
    )
    doc.add_page_break()


def page_brand_v3(doc: Document) -> None:
    add_page_header(doc, "01", "客户真正会记住什么", "参数会被忘掉，但身份感、价值观和未来可能性会被记住。")
    add_callout(
        doc,
        "品牌层主张",
        "智慧大脑不卖“更多功能”，而是帮助组织方把学生从“被评分的人”变成“能看见下一次成长路径的人”。",
        fill="EAF3FF",
        color=BLUE,
    )
    add_card_grid(doc, [
        ("客户看到的问题", "比赛、答辩和训练营往往停在一次评分，反馈难复盘，成长难证明。", BLUE),
        ("智慧大脑给出的答案", "把路演评审、AI 复盘、21 天训练、考试测评和能力档案连成一个闭环。", GREEN),
        ("组织方记住的价值", "不是多买一个系统，而是建立一套可持续复用的创新人才成长机制。", GOLD),
        ("学生记住的价值", "不只知道得了多少分，还知道下一轮该改什么、怎么练、练完如何被看见。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_table(
        doc,
        ["不该主讲", "应该主讲"],
        [
            ["技术栈、接口、数据库、后台菜单。", "评审如何留下证据，反馈如何转成行动。"],
            ["“我们有很多功能”。", "一次路演如何变成一套持续成长机制。"],
        ],
        [3.2, 3.15],
    )
    doc.add_page_break()


def page_flow_v3(doc: Document) -> None:
    add_page_header(doc, "03", "从路演现场到成长档案", "客户不需要记住所有模块，只要看懂这条路径。")
    add_table(
        doc,
        ["阶段", "用户看见什么", "平台沉淀什么"],
        [
            ["赛前", "项目材料、讲稿、评分点被提前组织。", "可训练的表达基础。"],
            ["赛中", "路演过程、评分依据、评委意见被记录。", "可追溯的评审证据。"],
            ["赛后", "分数、问题、AI 画像和建议被汇总。", "可执行的训练任务。"],
            ["训练", "21 天课程、每日任务、考试测评持续推进。", "可验证的成长过程。"],
            ["沉淀", "雷达图、趋势、证书、评级标签形成档案。", "可复用的组织资产。"],
        ],
        [1.05, 2.75, 2.55],
        first_col=True,
    )
    add_card_grid(doc, [
        ("赛前准备", "PPT 材料、素材证据、讲稿训练和评分点覆盖，让项目先具备表达能力。", BLUE),
        ("在线评审", "会议、演示、专家评分、团队路演评分和录制留痕，让过程可组织可复核。", GREEN),
        ("AI 复盘", "将转写、表达、现场呈现和评分结果收束为报告、问题和训练建议。", GOLD),
        ("问题闭环", "把扣分项和评委意见变成任务，而不是停留在一次性反馈。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "融合边界", "21 天训练当前作为融合规划模块表达，不能写成已经完全融合上线；宣传册里用“可接入、可承接、规划模块”保持可信。")
    doc.add_page_break()


def page_training_v3(doc: Document) -> None:
    add_page_header(doc, "04", "21 天训练让反馈落地", "把抽象建议拆成每天可完成、可记录、可验证的训练动作。")
    add_image(doc, ASSET_DIR / "ai_21day_journey.png", width=6.85, before=2, after=6)
    add_card_grid(doc, [
        ("理解问题", "查看复盘报告，明确短板与训练目标。", BLUE),
        ("补齐基础", "课程学习、资源阅读、每日任务提交。", GREEN),
        ("强化表达", "讲稿训练、打字速度、演示节奏优化。", GOLD),
        ("验证成果", "在线考试、团队评分、能力档案更新。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "买方语言", "训练平台不应被讲成课程后台，而要被讲成“把评审反馈落到每天行动，并把成长结果重新验证”的承接系统。")
    doc.add_page_break()


def page_verify_v3(doc: Document) -> None:
    add_page_header(doc, "05", "训练结果要被验证", "成长不是完成率，而是能力是否真的发生变化。")
    add_card_grid(doc, [
        ("在线考试", "考试创建、发布、分配、作答、提交和成绩查看，让训练有验收。", BLUE),
        ("题库练习", "单选、多选、判断、填空、编程题，构成可持续练习体系。", GREEN),
        ("批改反馈", "客观题自动判分，主观题和编程题支持人工批改与评语。", GOLD),
        ("AI 分析", "基于作答数据生成成绩分析和个性化学习建议，依赖系统 AI 配置启用。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_table(
        doc,
        ["验证对象", "验证方式", "形成资产"],
        [
            ["表达能力", "讲稿训练、路演评分、评委反馈。", "表达短板与改进记录。"],
            ["学习投入", "课程进度、每日任务、资源阅读。", "可追踪的训练过程。"],
            ["知识掌握", "在线考试、题库练习、成绩分析。", "可比对的测评结果。"],
            ["组织成效", "团队评分、证书、能力档案。", "可复用的培养数据。"],
        ],
        [1.35, 2.5, 2.5],
        first_col=True,
    )
    doc.add_page_break()


def page_profile_v3(doc: Document) -> None:
    add_page_header(doc, "06", "能力档案让成长被看见", "学生、教师和组织方看到的不再只是一个分数，而是一条成长轨迹。")
    add_image(doc, ASSET_DIR / "ai_dashboard.png", width=6.85, before=2, after=6)
    add_card_grid(doc, [
        ("六维雷达", "解决问题、代码、沟通、团队协作、演讲、创意。", BLUE),
        ("评分趋势", "路演总分与历史评分变化，观察团队表现。", GREEN),
        ("证书认证", "证书模板、颁发、查看与下载，沉淀学习成果。", GOLD),
        ("组织运营", "用户、部门、课程、资源、权限和日志统一管理。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_scenarios_trust_v3(doc: Document) -> None:
    add_page_header(doc, "07", "谁会需要智慧大脑", "同一套成长闭环，可以服务赛事、课程、园区和训练营。")
    add_table(
        doc,
        ["场景", "使用方式", "用户记住什么"],
        [
            ["创新创业大赛", "赛前准备、在线路演、专家评审、赛后训练。", "比赛不是终点，而是下一轮成长的起点。"],
            ["课程答辩", "教师创建任务，学生演示，系统沉淀评分与问题。", "教学评价可追溯，改进动作可跟踪。"],
            ["校内训练营", "多轮路演结合 21 天训练和考试测评。", "训练过程被看见，能力变化被证明。"],
            ["园区/孵化器", "项目材料打磨、专家评审和成长档案。", "项目服务从一次辅导升级为持续陪跑。"],
        ],
        [1.35, 2.55, 2.45],
        first_col=True,
    )
    add_card_grid(doc, [
        ("数据可信", "学习记录、评分记录、考试成绩、资源和证书统一沉淀，可追溯。", BLUE),
        ("组织可信", "支持角色权限、数据权限、管理日志和组织账号集成。", GREEN),
        ("内容可信", "支持防刷课、跑马灯水印、防录屏提醒等学习质量保障机制。", GOLD),
        ("交付可信", "支持多端访问、对象存储、私有化部署和系统配置能力。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    doc.add_page_break()


def page_landing_v3(doc: Document) -> None:
    add_page_header(doc, "08", "如何落地", "先验证一场真实评审，再把复盘结果接入训练模块。")
    add_table(
        doc,
        ["步骤", "建议动作", "阶段成果"],
        [
            ["1. 选场景", "选择一场创新创业比赛、课程答辩或训练营。", "明确角色、评分表和流程边界。"],
            ["2. 跑评审", "完成材料准备、在线路演、专家评分和复盘报告。", "验证评审组织与反馈质量。"],
            ["3. 接训练", "把复盘建议接入 21 天训练任务、课程和测评。", "验证从反馈到行动的闭环。"],
            ["4. 沉档案", "查看能力雷达、趋势、证书和组织数据。", "形成可复用的成长资产。"],
        ],
        [1.05, 3.0, 2.3],
        first_col=True,
    )
    add_callout(
        doc,
        "技术内容边界",
        "正文不展开 React、Spring Boot、MySQL、MinIO、Docker 等技术栈；这些应放在技术白皮书或部署文档中。宣传册只保留采购、部署、运营能理解的信任语言。",
    )
    add_card_grid(doc, [
        ("第一场试运行", "用真实路演验证流程、评分与复盘质量。", BLUE),
        ("第二阶段接训练", "把复盘建议转成 21 天任务和测评。", GREEN),
        ("第三阶段做运营", "沉淀模板、资源、数据和能力档案。", GOLD),
    ], cols=3, widths=[2.12, 2.12, 2.12])
    doc.add_page_break()


def page_close(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_a4_section(section)
    section.different_first_page_header_footer = True
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    for part in [section.header, section.footer, section.first_page_header, section.first_page_footer]:
        part.paragraphs[0].text = ""
    add_full_page_image(doc, ASSET_DIR / "closing_v2.png", width=7.03)


def build_doc() -> Document:
    make_cover()
    make_close()
    doc = new_doc()
    add_full_page_image(doc, ASSET_DIR / "cover_v2.png", width=7.03)
    doc.add_page_break()
    page_brand_v3(doc)
    page_loop(doc)
    page_flow_v3(doc)
    page_training_v3(doc)
    page_verify_v3(doc)
    page_profile_v3(doc)
    page_scenarios_trust_v3(doc)
    page_landing_v3(doc)
    page_close(doc)
    doc.core_properties.title = "智慧大脑产品宣传手册 - 视觉增强版V3"
    doc.core_properties.subject = "创新创业路演评审与成长训练平台"
    doc.core_properties.author = "Codex"
    doc.core_properties.created = datetime.now()
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
