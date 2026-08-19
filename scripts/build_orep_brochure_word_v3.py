from __future__ import annotations

from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_orep_brochure import (
    BLUE,
    CYAN,
    GOLD,
    GREEN,
    INK,
    LIGHT_BLUE,
    MUTED,
    NAVY,
    OUT_DIR,
    add_bullets,
    add_callout,
    add_heading,
    add_header_footer,
    add_p,
    add_section_divider,
    add_table,
    rgb,
    set_cell_border,
    set_cell_margins,
    set_cell_shading,
    set_page,
    set_run,
    set_table_width,
    style_table,
)


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = OUT_DIR / "word_v3_assets"
OUT_DOCX = OUT_DIR / "OREP在线路演评审平台_产品宣传手册_Word视觉增强版.docx"

FONT_REG = "/System/Library/Fonts/STHeiti Light.ttc"
FONT_BOLD = "/System/Library/Fonts/STHeiti Medium.ttc"

P = {
    "navy": "#061B44",
    "blue": "#0B4DBA",
    "bright": "#1B7CFF",
    "cyan": "#00A6FF",
    "ink": "#122033",
    "muted": "#5E6B7E",
    "soft": "#F4F7FB",
    "line": "#D7E2F0",
    "pale": "#EAF3FF",
    "green": "#0E7C70",
    "green_pale": "#EAFBF7",
    "amber": "#B7791F",
    "amber_pale": "#FFF7E6",
    "red": "#B42318",
    "white": "#FFFFFF",
}


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def draw_text(draw, xy, text, size=28, fill="#122033", bold=False, max_chars=None, line_gap=8):
    x, y = xy
    f = pil_font(size, bold)
    lines = []
    for para in text.split("\n"):
        if max_chars:
            while len(para) > max_chars:
                cut = max_chars
                for mark in "，。；、 ":
                    pos = para.rfind(mark, 0, max_chars + 1)
                    if pos > max_chars * 0.55:
                        cut = pos + 1
                        break
                lines.append(para[:cut])
                para = para[cut:]
        lines.append(para)
    for line in lines:
        draw.text((x, y), line, font=f, fill=fill)
        box = draw.textbbox((x, y), line or "口", font=f)
        y += (box[3] - box[1]) + line_gap
    return y


def rounded(draw, box, radius=18, fill="#FFFFFF", outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_dashboard(path: Path):
    w, h = 1400, 760
    img = Image.new("RGB", (w, h), "#F4F7FB")
    d = ImageDraw.Draw(img)
    rounded(d, (40, 36, w - 40, h - 36), 28, "#FFFFFF", "#D7E2F0", 3)
    rounded(d, (40, 36, w - 40, 120), 28, P["navy"])
    d.rectangle((40, 86, w - 40, 120), fill=P["navy"])
    draw_text(d, (78, 66), "OREP 路演评审工作台", 30, "#FFFFFF", True)
    for i, c in enumerate([P["cyan"], "#6E86B5", "#6E86B5", "#6E86B5"]):
        rounded(d, (760 + i * 110, 74, 840 + i * 110, 90), 8, c)
    rounded(d, (90, 164, 730, 570), 24, "#EAF3FF", "#BBD3F7", 3)
    rounded(d, (120, 200, 398, 392), 18, "#D7E8FF")
    d.ellipse((155, 248, 245, 338), fill=P["blue"])
    draw_text(d, (278, 272), "路演中", 38, P["navy"], True)
    rounded(d, (430, 208, 680, 300), 16, "#FFFFFF", "#D7E2F0", 2)
    draw_text(d, (458, 226), "评分进度 4/5", 25, P["blue"], True)
    rounded(d, (458, 270, 642, 286), 8, "#EAF3FF")
    rounded(d, (458, 270, 596, 286), 8, P["blue"])
    rounded(d, (430, 330, 680, 520), 16, "#EAFBF7", "#C9EEE7", 2)
    draw_text(d, (458, 356), "AI 复盘生成中", 25, P["green"], True)
    for i in range(5):
        rounded(d, (458, 410 + i * 22, 635 - (i % 2) * 42, 420 + i * 22), 5, "#A7DED3")
    for i, (title, body, color) in enumerate([
        ("会议协同", "参会、演示、互动", P["blue"]),
        ("专家评分", "分数、评论、依据", P["green"]),
        ("AI 画像", "诊断、证据、训练", P["amber"]),
    ]):
        x = 800
        y = 164 + i * 136
        rounded(d, (x, y, w - 96, y + 105), 18, "#FFFFFF", "#D7E2F0", 2)
        d.rectangle((x, y, x + 9, y + 105), fill=color)
        draw_text(d, (x + 32, y + 22), title, 29, color, True)
        draw_text(d, (x + 32, y + 60), body, 23, P["muted"])
    rounded(d, (90, 610, w - 90, 670), 18, P["pale"])
    for i, txt in enumerate(["赛前准备", "赛中评审", "赛后复盘", "持续运营"]):
        draw_text(d, (145 + i * 300, 626), txt, 24, P["blue"], True)
    img.save(path)


def make_cover(path: Path):
    w, h = 1600, 2264
    img = Image.new("RGB", (w, h), P["navy"])
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, w, h), fill=P["navy"])
    d.rectangle((0, 0, 150, h), fill="#08265F")
    d.rectangle((150, 0, 190, h), fill=P["bright"])
    rounded(d, (245, 190, 735, 250), 18, "#092B68", "#1B7CFF", 2)
    draw_text(d, (276, 205), "OREP PRODUCT BROCHURE", 30, P["cyan"], True)
    draw_text(d, (245, 370), "在线路演评审平台", 84, "#FFFFFF", True)
    draw_text(d, (250, 510), "面向高校、赛事与创新创业项目评审的一体化数字评审与成长平台", 39, "#DCEBFF", False, max_chars=26, line_gap=14)
    draw_text(d, (250, 690), "把路演评审从一次活动，升级为可组织、可追溯、可复盘、可持续运营的数字化体系。", 36, "#BFD8FF", False, max_chars=27, line_gap=16)

    rounded(d, (245, 920, 1420, 1668), 34, "#F4F7FB", "#2F66C9", 3)
    rounded(d, (310, 1000, 1355, 1100), 30, "#FFFFFF", "#C9D9F2", 3)
    draw_text(d, (350, 1027), "OREP 路演评审工作台", 32, P["navy"], True)
    for i, c in enumerate([P["cyan"], "#9DB3D5", "#9DB3D5", "#9DB3D5"]):
        rounded(d, (1035 + i * 72, 1045, 1086 + i * 72, 1063), 8, c)
    rounded(d, (310, 1140, 815, 1535), 28, "#EAF3FF", "#B7CEF0", 3)
    rounded(d, (365, 1205, 582, 1380), 20, "#D7E8FF")
    d.ellipse((414, 1252, 510, 1348), fill=P["blue"])
    draw_text(d, (610, 1252), "路演中", 44, P["navy"], True)
    rounded(d, (610, 1330, 760, 1395), 18, "#FFFFFF", "#D7E2F0", 2)
    draw_text(d, (630, 1348), "评分 4/5", 25, P["blue"], True)
    for i, (title, body, color) in enumerate([
        ("会议协同", "参会、演示、互动", P["blue"]),
        ("专家评分", "分数、评论、依据", P["green"]),
        ("AI 复盘", "诊断、证据、训练", P["amber"]),
    ]):
        y = 1140 + i * 132
        rounded(d, (875, y, 1300, y + 92), 22, "#FFFFFF", "#D7E2F0", 2)
        d.rectangle((875, y, 888, y + 92), fill=color)
        draw_text(d, (915, y + 20), title, 30, color, True)
        draw_text(d, (1070, y + 25), body, 24, P["muted"])
    rounded(d, (310, 1572, 1355, 1628), 18, P["pale"])
    for i, txt in enumerate(["赛前准备", "赛中评审", "赛后复盘", "持续运营"]):
        draw_text(d, (380 + i * 245, 1587), txt, 25, P["blue"], True)

    for i, (label, copy) in enumerate([
        ("组织方", "少靠人工串联"),
        ("评委", "评分反馈可追溯"),
        ("参赛团队", "知道下一轮怎么改"),
        ("IT 团队", "私有部署与数据治理"),
    ]):
        x = 245 + i * 295
        rounded(d, (x, 1810, x + 250, 1955), 26, "#092B68", "#224F9C", 2)
        draw_text(d, (x + 28, 1845), label, 30, "#FFFFFF", True)
        draw_text(d, (x + 28, 1892), copy, 24, "#BFD8FF", False, max_chars=8)
    draw_text(d, (245, 2115), "正式宣传版  |  2026", 28, "#8EABD8", True)
    img.save(path)


def make_loop(path: Path):
    w, h = 1400, 760
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    rounded(d, (40, 40, w - 40, h - 40), 28, "#F8FBFF", "#D7E2F0", 3)
    cx, cy = 700, 380
    d.ellipse((cx - 135, cy - 135, cx + 135, cy + 135), fill=P["pale"], outline="#C5D7EE", width=3)
    d.ellipse((cx - 95, cy - 95, cx + 95, cy + 95), fill=P["navy"])
    draw_text(d, (cx - 58, cy - 42), "OREP\n评审闭环", 35, "#FFFFFF", True, line_gap=10)
    cards = [
        ("赛前准备", "PPT / 素材 / 讲稿", 120, 135, P["blue"]),
        ("赛中评审", "会议 / 评分 / 录制", 920, 135, P["green"]),
        ("赛后复盘", "报告 / 画像 / 问题", 920, 500, P["amber"]),
        ("持续运营", "资源 / 数据 / 治理", 120, 500, P["blue"]),
    ]
    for title, body, x, y, color in cards:
        rounded(d, (x, y, x + 330, y + 130), 22, "#FFFFFF", "#D7E2F0", 3)
        d.rectangle((x, y, x + 10, y + 130), fill=color)
        draw_text(d, (x + 36, y + 30), title, 34, color, True)
        draw_text(d, (x + 36, y + 78), body, 24, P["muted"])
        d.line((x + 165, y + 65, cx, cy), fill="#BFD0E7", width=4)
    img.save(path)


def make_journey(path: Path):
    w, h = 1400, 840
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    roles = [
        ("组织方", ["创建会议", "配置评分", "发布资源", "汇总结果"]),
        ("参赛团队", ["准备材料", "进入会议", "完成路演", "查看反馈"]),
        ("评委", ["进入会议", "观看路演", "提交评分", "补充意见"]),
        ("管理员", ["维护角色", "管理资源", "监控数据", "沉淀历史"]),
    ]
    for i, (role, acts) in enumerate(roles):
        y = 52 + i * 186
        rounded(d, (40, y, w - 40, y + 132), 22, "#F8FBFF" if i % 2 == 0 else "#FFFFFF", "#D7E2F0", 3)
        draw_text(d, (78, y + 48), role, 34, P["navy"], True)
        for j, act in enumerate(acts):
            x = 310 + j * 230
            rounded(d, (x, y + 38, x + 155, y + 92), 14, "#FFFFFF", "#D7E2F0", 2)
            draw_text(d, (x + 26, y + 54), act, 24, P["blue"] if j == 2 else P["ink"], True)
            if j < 3:
                d.line((x + 164, y + 65, x + 205, y + 65), fill=P["bright"], width=4)
                d.polygon([(x + 205, y + 56), (x + 223, y + 65), (x + 205, y + 74)], fill=P["bright"])
    img.save(path)


def make_trust(path: Path):
    w, h = 1400, 760
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    items = [
        ("流程可信", "会议、评分、问题、报告和历史记录统一沉淀。", P["blue"]),
        ("证据可信", "反馈来自转写、表达节奏、现场呈现、材料证据和评分结果。", P["green"]),
        ("数据可信", "支持私有化部署和组织内数据治理。", P["amber"]),
        ("交付可信", "模块化服务、统一入口、文件管理、数据存储和后台监控。", P["blue"]),
    ]
    for i, (title, body, color) in enumerate(items):
        x = 70 + (i % 2) * 650
        y = 70 + (i // 2) * 210
        rounded(d, (x, y, x + 560, y + 150), 24, "#FFFFFF", "#D7E2F0", 3)
        d.rectangle((x, y, x + 10, y + 150), fill=color)
        draw_text(d, (x + 42, y + 30), title, 34, color, True)
        draw_text(d, (x + 42, y + 80), body, 25, P["muted"], max_chars=22, line_gap=8)
    rounded(d, (70, 520, w - 70, 690), 24, P["green_pale"], "#B8E8DE", 3)
    draw_text(d, (110, 555), "技术内容边界", 34, P["green"], True)
    draw_text(d, (110, 610), "宣传手册只说明部署形态、数据控制、实时协同和 AI 分析能力；具体框架、接口、端口、数据库表和配置方式应放在技术白皮书或部署文档。", 26, P["ink"], max_chars=44)
    img.save(path)


def make_closing(path: Path):
    w, h = 1600, 2264
    img = Image.new("RGB", (w, h), P["navy"])
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, w, h), fill=P["navy"])
    d.rectangle((0, h - 390, w, h), fill="#071735")
    draw_text(d, (190, 265), "从活动工具", 82, "#FFFFFF", True)
    draw_text(d, (190, 395), "到数字评审基础设施", 82, "#FFFFFF", True)
    draw_text(d, (195, 620), "OREP 的价值不止是线上化，而是把路演评审的组织、证据、反馈和训练变成可运营资产。", 38, "#DCEBFF", False, max_chars=28, line_gap=16)
    labels = [
        ("可组织", "统一流程和角色"),
        ("可追溯", "评分、评论、录制留痕"),
        ("可复盘", "报告、问题、AI 画像"),
        ("可运营", "资源、数据、历史沉淀"),
    ]
    for i, (title, body) in enumerate(labels):
        x = 190 + (i % 2) * 590
        y = 920 + (i // 2) * 245
        rounded(d, (x, y, x + 500, y + 155), 28, "#092B68", "#244F96", 3)
        draw_text(d, (x + 38, y + 32), title, 40, P["cyan"], True)
        draw_text(d, (x + 38, y + 92), body, 28, "#DCEBFF")
    rounded(d, (190, 1585, 1410, 1780), 34, "#FFFFFF", "#C7D7EF", 3)
    draw_text(d, (235, 1630), "建议下一步", 38, P["blue"], True)
    draw_text(d, (235, 1690), "选择一个真实评审场景，完成 1 场试运行，验证流程、反馈质量和组织效率。", 31, P["ink"], False, max_chars=34, line_gap=12)
    draw_text(d, (190, 2050), "适用于：高校创新创业学院、赛事承办方、园区/孵化器、课程答辩与项目训练营", 30, "#AFC8EC", False)
    img.save(path)


def create_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_cover(ASSET_DIR / "cover.png")
    make_dashboard(ASSET_DIR / "dashboard.png")
    make_loop(ASSET_DIR / "loop.png")
    make_journey(ASSET_DIR / "journey.png")
    make_trust(ASSET_DIR / "trust.png")
    make_closing(ASSET_DIR / "closing.png")


def configure_a4_section(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)


def set_a4_page(doc: Document):
    configure_a4_section(doc.sections[0])


def add_image(doc: Document, path: Path, width=7.0, before=4, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))


def add_full_page_image(doc: Document, path: Path, width=7.03):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))


def add_card_grid(doc: Document, cards: list[tuple[str, str, str]], cols=2, widths=None):
    if widths is None:
        widths = [3.25] * cols
    rows = (len(cards) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    set_table_width(table, widths)
    for idx in range(rows * cols):
        r = idx // cols
        c = idx % cols
        cell = table.cell(r, c)
        set_cell_margins(cell, top=145, bottom=145, start=160, end=160)
        set_cell_shading(cell, "FFFFFF")
        set_cell_border(cell, color="C9D9F2", size="8")
        if idx >= len(cards):
            cell.text = ""
            set_cell_border(cell, color="FFFFFF", size="1")
            continue
        title, body, color = cards[idx]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        rr = p.add_run(title)
        set_run(rr, size=11.5, color=color, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_after = Pt(0)
        rr2 = p2.add_run(body)
        set_run(rr2, size=9.7, color=INK)
    add_p(doc, "", after=4)
    return table


def add_dark_cover(doc: Document):
    add_full_page_image(doc, ASSET_DIR / "cover.png", width=7.03)
    doc.add_page_break()


def page_why(doc: Document):
    add_section_divider(doc, "01", "为什么需要 OREP", "路演评审的难点，不只是把会议搬到线上，而是让评审过程真正可管理、可解释、可复用。")
    add_card_grid(doc, [
        ("组织分散", "会议、材料、评分、反馈和报告分散在多个工具，活动组织依赖人工协调。OREP 将赛前、赛中、赛后和运营管理放在同一条链路。", BLUE),
        ("评分难追溯", "分数之外缺少扣分依据、问题记录和可归档材料，复核成本高。OREP 形成评分记录、问题闭环、录制留痕和报告输出。", GREEN),
        ("反馈不成体系", "学生只知道结果，不知道表达、材料、逻辑和现场呈现具体怎么改。OREP 把 AI 画像、表达诊断和训练建议转成行动。", GOLD),
        ("经验难复用", "每次比赛重新组织，模板、资源、历史项目和数据难沉淀。OREP 建立资源库、历史任务、数据监控和后台治理能力。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "读者应当记住", "OREP 的核心不是“多一个路演工具”，而是让组织方拥有一套数字化评审基础设施。")
    doc.add_page_break()


def page_loop(doc: Document):
    add_section_divider(doc, "02", "产品总览：四段闭环", "OREP 把准备、评审、复盘和运营放进同一个产品闭环，让所有参与方在统一流程中协作。")
    add_image(doc, ASSET_DIR / "loop.png", width=7.0)
    add_table(
        doc,
        ["阶段", "用户动作", "平台交付"],
        [
            ["赛前准备", "学生准备项目材料，教师/组织方提供模板和要求。", "PPT 生成、素材证据、讲稿训练、资源中心。"],
            ["赛中评审", "团队在线路演，评委参与会议并完成评分。", "在线会议、专家评分、聊天互动、录制留痕。"],
            ["赛后复盘", "组织方汇总结果，学生查看反馈并进入下一轮优化。", "评分报告、问题闭环、AI 能力画像、训练建议。"],
            ["持续运营", "学校/赛事/园区沉淀资源、数据、历史项目和管理经验。", "后台治理、数据监控、资源管理、私有化部署。"],
        ],
        [1.2, 2.65, 2.5],
        first_col=True,
    )
    doc.add_page_break()


def page_journey(doc: Document):
    add_section_divider(doc, "03", "一场路演如何运转", "用真实工作流解释产品，而不是让读者在功能清单里寻找答案。")
    add_image(doc, ASSET_DIR / "journey.png", width=7.0)
    add_callout(doc, "流程图的意义", "每个角色都能在同一条评审链路里完成自己的工作，并让过程沉淀为数据、证据和改进动作。")
    doc.add_page_break()


def page_prepare(doc: Document):
    add_section_divider(doc, "04", "赛前准备：让材料先具备得分能力", "路演质量不是从上台开始，而是从材料结构、证据完整性和讲稿准备开始。")
    add_card_grid(doc, [
        ("PPT 智能生成", "将项目材料组织成路演结构，生成可预览、可编辑、可下载的材料初稿。", BLUE),
        ("评分点覆盖", "让团队看到哪些评分点已覆盖，哪些证据、页面或叙事还需要补强。", GREEN),
        ("素材证据管理", "把政策、截图、设备照片、数据图等材料与页面、实操步骤和评分点绑定。", GOLD),
        ("讲稿训练", "从页面内容生成讲稿和训练文本，让彩排不只依赖临场发挥。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "对参赛团队的吸引点", "OREP 不是只帮你做一份 PPT，而是帮你把 PPT 做成能被评委理解、能支撑评分、能用于训练的路演材料。")
    doc.add_page_break()


def page_review(doc: Document):
    add_section_divider(doc, "05", "赛中评审：有序、留痕、可复核", "评审现场需要稳定协同，也需要把关键过程转化为后续复盘依据。")
    add_table(
        doc,
        ["能力", "用户体验", "管理价值"],
        [
            ["在线会议", "团队、评委和组织方通过统一入口参与路演。", "减少线下排期和跨工具组织成本。"],
            ["专家评分", "评委按评分项提交分数、评论和判断依据。", "评分口径更集中，结果更容易汇总。"],
            ["互动协同", "会议内支持沟通、状态流转和参会人管理。", "现场管理更可控。"],
            ["录制留痕", "路演过程可保存并用于复盘。", "为争议复核、AI 分析和教学反馈提供素材基础。"],
        ],
        [1.35, 2.5, 2.5],
        first_col=True,
    )
    add_callout(doc, "不把会议工具当终点", "会议只是评审现场，OREP 更重要的是把现场行为转化为评分、证据、报告和成长反馈。")
    doc.add_page_break()


def page_after(doc: Document):
    add_section_divider(doc, "06", "赛后复盘：让分数变成改进动作", "优秀的评审平台不只给结果，还要帮助学生和指导老师知道下一轮怎么提升。")
    add_card_grid(doc, [
        ("评分报告", "将评分结果、维度表现和关键问题整理为可归档材料，服务复核与教学反馈。", BLUE),
        ("问题闭环", "把扣分项、评论和待改进事项转化为可跟踪的问题记录。", GREEN),
        ("AI 能力画像", "结合转写、表达节奏、现场呈现和评分结果，形成描述性反馈。", GOLD),
        ("训练建议", "把表达、材料、演示和证据短板收束为下一轮可执行任务。", BLUE),
    ], cols=2, widths=[3.25, 3.25])
    add_callout(doc, "AI 的正确讲法", "不要只说“AI 赋能”。手册里应说明 AI 具体产出什么：转写、表达诊断、能力画像、证据链和训练建议。")
    doc.add_page_break()


def page_operation(doc: Document):
    add_section_divider(doc, "07", "持续运营：让一次比赛沉淀为长期资产", "当平台持续承载比赛、课程和训练营，组织方才能积累资源、经验和数据。")
    add_table(
        doc,
        ["运营对象", "平台支持", "沉淀资产"],
        [
            ["用户与角色", "维护管理员、教师、学生、评委、专家等角色。", "清晰的访问身份和组织权限。"],
            ["资源与模板", "统一管理 PPT 模板、评分表、交付文档和素材。", "可复用的赛事与课程资源库。"],
            ["会议与任务", "管理会议、评分、PPT 任务、历史记录和问题。", "完整的活动过程档案。"],
            ["数据监控", "查看在线状态、行为日志和平台使用情况。", "运营决策和服务改进依据。"],
        ],
        [1.4, 2.45, 2.5],
        first_col=True,
    )
    doc.add_page_break()


def page_scenarios(doc: Document):
    add_section_divider(doc, "08", "典型应用场景", "同一套评审闭环，可以服务多种路演、答辩和项目评估场景。")
    add_table(
        doc,
        ["场景", "使用方式", "吸引点"],
        [
            ["创新创业大赛", "赛前准备材料，赛中专家评分，赛后输出报告和训练建议。", "提升组织效率和评分规范性。"],
            ["课程答辩/项目验收", "教师创建会议，学生在线演示，平台留存评分与问题。", "让教学评价可追溯、可复盘。"],
            ["园区/孵化器路演", "为项目提供材料打磨、专家评审和复盘服务。", "提升项目展示质量和投融资准备度。"],
            ["校内训练营", "多轮路演训练，每轮保留评分、画像和改进任务。", "把比赛准备变成持续训练机制。"],
            ["远程专家评审", "跨地域专家进入会议并完成标准化评分。", "降低专家组织成本，扩大评审覆盖。"],
        ],
        [1.4, 2.65, 2.3],
        first_col=True,
    )
    doc.add_page_break()


def page_trust(doc: Document):
    add_section_divider(doc, "09", "为什么可信", "技术信息在宣传手册中应转化为采购和部署能理解的信任语言。")
    add_image(doc, ASSET_DIR / "trust.png", width=7.0)
    doc.add_page_break()


def page_get_started(doc: Document):
    add_section_divider(doc, "10", "如何落地", "从小规模试运行开始，让组织方先验证流程，再扩大到赛事、课程和训练营。")
    add_table(
        doc,
        ["步骤", "建议动作", "阶段成果"],
        [
            ["1. 明确场景", "选择一类先行场景，如创新创业大赛初赛或课程答辩。", "确定评分表、参与角色和流程边界。"],
            ["2. 配置资源", "准备模板、评分项、会议安排、账号角色和材料要求。", "形成可执行的评审空间。"],
            ["3. 试运行", "用一场小规模路演验证会议、评分、报告和复盘。", "发现流程问题并优化组织方式。"],
            ["4. 正式运营", "推广到更多项目、班级、赛事或园区活动。", "沉淀资源、历史数据和持续训练机制。"],
        ],
        [1.05, 3.05, 2.25],
        first_col=True,
    )
    add_callout(doc, "建议下一步", "选择一个真实评审场景，先做 1 场试运行，验证流程、反馈和组织效率。")


def page_close(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_a4_section(section)
    section.different_first_page_header_footer = True
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""
    add_full_page_image(doc, ASSET_DIR / "closing.png", width=7.03)


def build_doc() -> Document:
    create_assets()
    doc = Document()
    set_a4_page(doc)
    doc.sections[0].different_first_page_header_footer = True
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["List Bullet"].font.name = "Microsoft YaHei"
    styles["List Bullet"].font.size = Pt(10.5)
    add_header_footer(doc.sections[0], "OREP Product Brochure")
    add_dark_cover(doc)
    page_why(doc)
    page_loop(doc)
    page_journey(doc)
    page_prepare(doc)
    page_review(doc)
    page_after(doc)
    page_operation(doc)
    page_scenarios(doc)
    page_trust(doc)
    page_get_started(doc)
    page_close(doc)
    doc.core_properties.title = "OREP 在线路演评审平台产品宣传手册 - Word视觉增强版"
    doc.core_properties.subject = "参考 Figma 样式与配色重制"
    doc.core_properties.author = "Codex"
    doc.core_properties.created = datetime.now()
    return doc


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
