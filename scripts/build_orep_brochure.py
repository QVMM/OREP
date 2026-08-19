from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "product_brochure"
OUT_DOCX = OUT_DIR / "OREP在线路演评审平台_产品宣传手册.docx"


BLUE = "0B3D91"
CYAN = "1A73E8"
NAVY = "0B1F4D"
INK = "172033"
MUTED = "5F6B7A"
LIGHT_BLUE = "EAF3FF"
PALE = "F5F8FC"
GRID = "D8E1EC"
GREEN = "0F766E"
GOLD = "B7791F"
RED = "B42318"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths_in):
            row.cells[i].width = Inches(width)
            tc_pr = row.cells[i]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_run(run, size=11, color=INK, bold=False, font="Microsoft YaHei") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def add_p(doc, text="", size=11, color=INK, bold=False, before=0, after=6, align=None, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    color = BLUE if level <= 2 else NAVY
    size = {1: 18, 2: 14, 3: 12}.get(level, 12)
    before = {1: 14, 2: 10, 3: 6}.get(level, 6)
    after = {1: 8, 2: 6, 3: 4}.get(level, 4)
    p = add_p(doc, text, size=size, color=color, bold=True, before=before, after=after)
    return p


def add_bullets(doc, items: list[str], level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28 + level * 0.18)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_run(r, size=10.5)


def style_table(table, header=True, first_col=False) -> None:
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            fill = WHITE
            if header and row_idx == 0:
                fill = BLUE
            elif first_col and col_idx == 0:
                fill = LIGHT_BLUE
            elif row_idx % 2 == 0:
                fill = "FAFCFF"
            set_cell_shading(cell, fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.12
                for r in p.runs:
                    if header and row_idx == 0:
                        set_run(r, size=9.5, color=WHITE, bold=True)
                    elif first_col and col_idx == 0:
                        set_run(r, size=9.5, color=BLUE, bold=True)
                    else:
                        set_run(r, size=9.5, color=INK)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float], first_col=False):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].text = text
    style_table(table, header=True, first_col=first_col)
    add_p(doc, "", after=4)
    return table


def add_callout(doc, title: str, body: str, fill=LIGHT_BLUE, color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.35])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="BBD4F6", size="8")
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run(r, size=11.5, color=color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.16
    r2 = p2.add_run(body)
    set_run(r2, size=10.5, color=INK)
    add_p(doc, "", after=4)


def set_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def add_header_footer(section, title="OREP 在线路演评审平台") -> None:
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run(title)
    set_run(r, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = fp.add_run("产品宣传手册 | 2026")
    set_run(r2, size=8, color=MUTED)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    add_header_footer(section, "OREP Product Brochure")
    add_p(doc, "OREP PRODUCT BROCHURE", size=10, color=CYAN, bold=True, before=16, after=22)
    p = add_p(doc, "在线路演评审平台", size=30, color=NAVY, bold=True, before=0, after=8)
    p.paragraph_format.line_spacing = 1.05
    add_p(doc, "面向高校、赛事、园区与创新创业项目评审的一体化数字评审基础设施", size=15, color=INK, before=0, after=18)
    add_callout(
        doc,
        "核心主张",
        "OREP 将在线会议、专家评分、AI 音视频分析、PPT 智能生成、讲稿训练、资源管理与后台治理整合为一套闭环平台，让路演评审从一次性活动升级为可组织、可追溯、可复盘、可持续提升的数字化能力。",
        fill=LIGHT_BLUE,
    )
    metrics = [
        ["赛前", "PPT 生成、素材证据、讲稿准备、模板资源"],
        ["赛中", "在线会议、实时协同、专家评分、录制留痕"],
        ["赛后", "评分报告、AI 画像、问题跟踪、训练建议"],
        ["运营", "后台治理、数据监控、资源管理、私有化部署"],
    ]
    add_table(doc, ["阶段", "平台能力"], metrics, [1.1, 5.25], first_col=True)
    add_p(doc, "版本：正式宣传版 | 生成日期：2026-06-10", size=9, color=MUTED, before=22, after=0)
    doc.add_page_break()


def add_section_divider(doc: Document, index: str, title: str, subtitle: str) -> None:
    add_p(doc, index, size=10, color=CYAN, bold=True, before=8, after=8)
    add_p(doc, title, size=24, color=NAVY, bold=True, before=0, after=8)
    add_p(doc, subtitle, size=12, color=MUTED, before=0, after=18)


def build_doc() -> Document:
    doc = Document()
    set_page(doc)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["List Bullet"].font.name = "Microsoft YaHei"
    styles["List Bullet"].font.size = Pt(10.5)

    add_cover(doc)

    add_section_divider(
        doc,
        "01",
        "行业挑战与建设目标",
        "路演评审正在从线下组织、单点评分，走向线上协同、全过程留痕和数据驱动的成长评价。",
    )
    add_heading(doc, "传统路演评审的四类痛点", 2)
    add_table(
        doc,
        ["痛点", "典型表现", "OREP 建设目标"],
        [
            ["组织成本高", "会议、资料、评委、学生与结果分散在多个工具中，活动组织依赖人工串联。", "用统一入口承载会议、资料、评分、录制与报告。"],
            ["评分过程难追溯", "评分依据、扣分原因、问题反馈和复盘证据难以沉淀。", "形成评分记录、问题跟踪、PDF 报告与 AI 证据链。"],
            ["赛后成长弱", "路演结束后只得到分数，难以知道表达、逻辑、材料与现场呈现如何改进。", "输出能力画像、表达训练、PPT 证据补强与行动计划。"],
            ["平台化运营不足", "模板资源、项目材料、历史产物和运维监控缺少统一管理。", "建设可持续运营的后台、资源库、PPT 管理和数据监控体系。"],
        ],
        [1.35, 2.55, 2.45],
        first_col=True,
    )
    add_callout(
        doc,
        "建设愿景",
        "以 OREP 为在线路演评审底座，帮助学校、赛事主办方和园区孵化机构建立标准化评审流程、可复用资源体系和可解释的能力成长档案。",
    )

    doc.add_page_break()
    add_section_divider(
        doc,
        "02",
        "产品总览",
        "OREP 不是单一会议工具，而是一套围绕路演评审全流程设计的产品矩阵。",
    )
    add_heading(doc, "一平台、双门户、五大能力域", 2)
    add_table(
        doc,
        ["能力域", "面向用户", "核心能力"],
        [
            ["在线会议协同", "参赛团队、评委、教师、组织方", "会议创建、会议码加入、LiveKit 音视频、聊天协作、屏幕共享、会议历史、录制回看。"],
            ["专家评分评审", "评委、教师、赛事管理员", "评分项获取、评分提交、重复提交防护、结果查看、问题生成、PDF 导出。"],
            ["AI 评分与画像", "学生、教师、评委、训练营导师", "ASR 转写、语音质量分析、视觉关键帧、五维评分、融合时间线、能力画像、训练建议。"],
            ["PPT 与讲稿智能生产", "参赛团队、指导老师", "问卷生成、模板导入、在线预览、质量报告、评分点覆盖、素材证据、讲稿编辑、文件下载。"],
            ["后台治理与交付", "学校/赛事/园区运营团队", "用户角色、会议管理、资源管理、PPT 任务管理、数据监控、Docker/Nginx 私有化部署。"],
        ],
        [1.45, 1.55, 3.35],
        first_col=True,
    )
    add_heading(doc, "角色与入口", 2)
    add_bullets(
        doc,
        [
            "用户端：覆盖登录注册、会议、AI 评分结果、PPT 生成/编辑、讲稿、资源中心、个人中心和历史记录。",
            "管理端：覆盖统计概览、用户管理、会议管理、评分查看、问题跟踪、数据分析、数据监控、资源管理和 PPT 任务管理。",
            "AI 服务端：提供评分、实时转写、视觉帧、能力画像、表达辅导、PPT Agent 与报告生成接口。",
        ]
    )

    doc.add_page_break()
    add_section_divider(
        doc,
        "03",
        "赛中评审协同",
        "以在线会议为评审现场，以标准化评分与问题跟踪沉淀评审过程。",
    )
    add_heading(doc, "在线会议与实时协作", 2)
    add_bullets(
        doc,
        [
            "支持会议创建、会议列表、会议详情、会议码加入、会议开始/结束、参会人管理与历史详情查看。",
            "集成 LiveKit/WebRTC 能力，为路演、答辩、评委评审和屏幕演示提供音视频协同基础。",
            "支持聊天消息、会议记录、录制状态、录制回看与流式播放，便于赛后复盘和证据留存。",
        ]
    )
    add_heading(doc, "评审评分与问题闭环", 2)
    add_table(
        doc,
        ["功能", "说明"],
        [
            ["评分表单", "按预设评分项提交分数和评论，适配比赛评分标准与评委工作流。"],
            ["结果统计", "查看会议评分结果、综合分与评分明细，为组织方形成统一结果口径。"],
            ["PDF 导出", "评分结果和问题列表可导出为正式报告材料，支撑归档和汇报。"],
            ["问题跟踪", "扣分项或评论可形成问题记录，支持未解决问题查看、处理和闭环管理。"],
        ],
        [1.55, 4.8],
        first_col=True,
    )
    add_callout(doc, "价值输出", "将“打完分就结束”的评审方式升级为“有依据、有记录、有问题、有复盘”的闭环评审流程。")

    doc.add_page_break()
    add_section_divider(
        doc,
        "04",
        "AI 评分与能力画像",
        "OREP 将语音、文本、视觉与现场节奏纳入分析，为参赛团队提供可解释的成长反馈。",
    )
    add_heading(doc, "AI 分析能力矩阵", 2)
    add_table(
        doc,
        ["能力", "数据来源", "输出结果"],
        [
            ["实时/离线转写", "音频流、上传音频、会议录制", "ASR 文本、分段内容、时间戳与转写事件。"],
            ["语音质量分析", "演讲音频", "语速、停顿、口头禅、表达流畅度和表达训练建议。"],
            ["视觉关键帧分析", "视频帧、屏幕/摄像头画面", "姿态、表情、眼神、屏幕类型、现场呈现证据。"],
            ["五维评分", "ASR 文本、语音特征、视觉证据", "内容质量、表达力、逻辑结构、技术深度、团队协作评分。"],
            ["融合时间线", "音频窗口与视觉帧对齐", "语音、视觉和综合表现走势，定位高光与风险片段。"],
            ["能力画像", "评分结果、转写片段、融合数据", "团队分工、表达风格、关键洞察、改进建议和训练清单。"],
        ],
        [1.45, 1.75, 3.15],
        first_col=True,
    )
    add_heading(doc, "三类 AI 辅导场景", 2)
    add_bullets(
        doc,
        [
            "表达节奏辅导：针对语速偏快、停顿过长、口头禅频繁等问题生成训练建议。",
            "路演呈现辅导：结合屏幕证据和评委理解成本，定位材料讲解与演示衔接问题。",
            "行动计划辅导：把五维评分、现场证据和材料缺口收束为下一轮可执行优化任务。",
        ]
    )

    doc.add_page_break()
    add_section_divider(
        doc,
        "05",
        "PPT 智能生成与路演准备",
        "从项目材料到路演 PPT，再到讲稿和评分点覆盖，OREP 支持赛前准备的结构化生产。",
    )
    add_heading(doc, "PPT Agent 全流程", 2)
    add_table(
        doc,
        ["阶段", "平台动作", "产物"],
        [
            ["输入", "收集问卷、项目文档、素材、政策截图、数据图表和用户模板。", "结构化需求、素材证据库。"],
            ["生成", "调用 PPT Agent 生成大纲、页面 HTML、预览、路演计划和可下载 PPT。", "PPT 初稿、页面预览、下载文件。"],
            ["质检", "生成评分点覆盖、质量报告、实操演示建议和缺图任务。", "风险矩阵、补强清单、页面级优化建议。"],
            ["编辑", "支持在线预览、PPT 编辑器、React/Konva 工作区、讲稿编辑和页面讲稿持久化。", "可迭代演示材料、讲稿和历史版本。"],
            ["交付", "支持下载、重导出、模板管理、资源中心和历史详情。", "正式 PPT、讲稿、素材包和过程记录。"],
        ],
        [1.05, 3.0, 2.3],
        first_col=True,
    )
    add_heading(doc, "面向比赛评分的材料建设", 2)
    add_bullets(
        doc,
        [
            "评分点覆盖矩阵帮助团队判断页面是否真正服务比赛评分表，避免只追求视觉效果。",
            "素材证据管理支持政策官方来源、系统截图、设备照片、数据图、文档材料和视频关键帧绑定。",
            "质量报告对页面视觉风险、评分缺口和实操问题进行分级，辅助赛前集中补强。",
        ]
    )

    doc.add_page_break()
    add_section_divider(
        doc,
        "06",
        "后台治理与运营管理",
        "以管理端支撑组织方的长期运营，而不是只支撑一次比赛活动。",
    )
    add_heading(doc, "管理后台能力", 2)
    add_table(
        doc,
        ["模块", "能力"],
        [
            ["统计概览", "集中查看平台关键数据与运营态势。"],
            ["用户管理", "维护账号、邮箱、租户字段、角色与访问身份，覆盖管理员、校级管理员、教师、学生、评审员和专家等角色。"],
            ["会议管理", "创建会议、维护会议号/密码/状态、启动或结束会议。"],
            ["评分查看", "查看评分结果和评分明细，为赛事归档和复核提供依据。"],
            ["问题跟踪", "管理评审反馈中的问题项，形成从发现到解决的处理闭环。"],
            ["资源管理", "上传、下载、预览和删除 PPT 资源，为参赛团队提供统一资料入口。"],
            ["PPT 管理", "管理 PPT 任务、上传资源和生成产物，支撑赛前材料治理。"],
            ["数据监控", "采集页面访问、用户在线心跳和行为日志，帮助运营团队掌握使用状态。"],
        ],
        [1.45, 4.9],
        first_col=True,
    )
    add_callout(doc, "运营价值", "管理端让组织方具备“可配置、可管理、可监控、可归档”的平台能力，适合学校、赛事承办方和园区孵化机构持续使用。")

    doc.add_page_break()
    add_section_divider(
        doc,
        "07",
        "技术架构与部署交付",
        "OREP 采用前后端分离与多服务架构，兼顾实时协同、AI 计算和私有化部署。",
    )
    add_heading(doc, "系统架构", 2)
    add_table(
        doc,
        ["层级", "组件", "作用"],
        [
            ["前端层", "Vue 3、Element Plus、用户端与管理端", "承载参赛、评审、管理、PPT 编辑和 AI 结果展示体验。"],
            ["业务后端", "Spring Boot、MyBatis-Plus、JWT、PDF 服务", "提供认证、用户、会议、评分、问题、资源、录制和统计等核心业务接口。"],
            ["AI 服务", "FastAPI、ASR、LLM、视频分析、PPT Agent", "提供智能评分、能力画像、辅导建议、报告生成和 PPT 生成能力。"],
            ["实时通信", "LiveKit、Socket.IO、WebSocket", "支撑音视频会议、聊天、转写流、PPT 任务进度和信令协同。"],
            ["数据与存储", "MySQL、Redis、上传目录、MinIO 服务接口", "沉淀业务数据、会话状态、文件资源和历史产物。"],
            ["部署层", "Docker Compose、Nginx、LiveKit 配置、录制 Bot", "支持统一入口、服务编排、反向代理、录制处理和服务器部署。"],
        ],
        [1.05, 2.0, 3.3],
        first_col=True,
    )
    add_heading(doc, "交付特性", 2)
    add_bullets(
        doc,
        [
            "支持本地开发环境和服务器 Docker 部署，适合校内或机构私有化落地。",
            "Nginx 统一入口反向代理前端、后端、AI 服务、信令与 LiveKit 相关服务。",
            "MySQL/Redis 支撑业务持久化与状态缓存，部署包包含数据库初始化脚本和升级脚本。",
            "录制 Bot 与 LiveKit Egress 能力为会议留痕、赛后复盘和 AI 分析提供素材基础。",
        ]
    )

    doc.add_page_break()
    add_section_divider(
        doc,
        "08",
        "典型应用场景",
        "围绕不同组织目标，OREP 可以作为比赛平台、训练平台和项目评审平台落地。",
    )
    add_table(
        doc,
        ["场景", "使用方式", "关键收益"],
        [
            ["创新创业大赛", "赛前提交材料与生成路演 PPT，赛中在线评审，赛后输出评分和问题报告。", "提升组织效率，统一评分口径，沉淀赛事数据。"],
            ["课程答辩与项目验收", "教师创建会议，学生在线演示，系统记录评分、问题和报告。", "减少线下排期压力，让过程可追溯。"],
            ["孵化器/园区路演", "为入驻项目提供路演训练、材料优化和专家评审。", "帮助项目提高表达质量和融资材料专业度。"],
            ["校内训练营", "通过 AI 评分、能力画像、讲稿训练和 PPT 质检形成多轮迭代。", "把一次路演变成持续提升机制。"],
            ["远程专家评审", "专家跨地域进入会议并完成评分，后台统一汇总。", "降低专家组织成本，提高评审覆盖范围。"],
        ],
        [1.45, 2.65, 2.25],
        first_col=True,
    )
    add_callout(
        doc,
        "适配对象",
        "高校创新创业学院、教务/实践教学部门、赛事承办单位、产业园区、孵化器、创投服务机构和项目训练营均可围绕自身流程配置使用。",
    )

    doc.add_page_break()
    add_section_divider(
        doc,
        "09",
        "客户价值",
        "OREP 的价值不止是线上化，而是把路演评审的组织、证据、反馈和训练变成可运营资产。",
    )
    add_table(
        doc,
        ["价值维度", "客户可感知收益"],
        [
            ["规范评审", "评分项、评分记录、问题反馈和 PDF 报告统一沉淀，降低评审口径差异。"],
            ["提升效率", "会议、资料、评分、录制、报告和后台管理在一个平台内完成，减少工具切换。"],
            ["增强透明", "AI 分析提供转写、语音、视觉、融合时间线等证据，帮助解释评分与改进建议。"],
            ["促进成长", "学生获得能力画像、表达训练、材料补强和行动计划，提升下一轮路演质量。"],
            ["持续运营", "资源库、模板、历史任务、数据监控和私有化部署帮助组织方长期复用。"],
        ],
        [1.45, 4.9],
        first_col=True,
    )
    add_heading(doc, "一句话总结", 2)
    add_callout(
        doc,
        "OREP 在线路演评审平台",
        "把路演评审从“组织一场活动”升级为“建设一套可持续运营的数字化评审与成长体系”。",
        fill="EAFBF8",
        color=GREEN,
    )

    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.core_properties.title = "OREP 在线路演评审平台产品宣传手册"
    doc.core_properties.subject = "正式产品宣传手册"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "基于 OREP 当前仓库功能生成"
    doc.core_properties.created = datetime.now()
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
