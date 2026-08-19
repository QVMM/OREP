from __future__ import annotations

import asyncio
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.routers.ppt_router import health
from app.services.ppt.agent_runtime import bootstrap_ppt_agent

bootstrap_ppt_agent()

from backend.api.endpoints.history import _job_to_history_item  # noqa: E402
from backend.api.schemas import JobStatus  # noqa: E402
from backend.generator.image2_editable_export import create_image2_image_pptx  # noqa: E402
from backend.generator.svg_to_pptx import create_pptx  # noqa: E402
from backend.orchestrator.roadshow_script_skill import (  # noqa: E402
    _scaled_script_targets,
    _script_token_budget,
)
from backend.orchestrator.strategist_agent import _design_spec_token_budget  # noqa: E402
from backend.session.manager import Job  # noqa: E402


def test_health_exposes_image2_capability_and_export_mode() -> None:
    payload = asyncio.run(health())

    assert payload["status"] == "ok"
    assert isinstance(payload["image2_configured"], bool)
    assert payload["image2_export_mode"] in {"image", "editable", "svg"}


def test_status_and_history_preserve_render_engine() -> None:
    status = JobStatus(status="complete", render_engine="svg")
    job = Job(
        id="job-svg",
        session_id="session-1",
        status="complete",
        style="roadshow",
        render_engine="svg",
    )

    history_item = _job_to_history_item(job, None, include_slide_count=False)

    assert status.render_engine == "svg"
    assert history_item.deck_type == "roadshow"
    assert history_item.render_engine == "svg"


def test_image2_image_export_creates_one_full_slide_picture_per_page(tmp_path: Path) -> None:
    project_dir = tmp_path / "image2-project"
    image_dir = project_dir / "images" / "image2"
    image_dir.mkdir(parents=True)
    for page, color in ((1, (232, 74, 28)), (2, (44, 48, 56))):
        image_path = image_dir / f"page_{page:02d}_generated.png"
        Image.new("RGB", (1600, 900), color).save(image_path)

    output_path = tmp_path / "image-output.pptx"
    result, warnings = create_image2_image_pptx(project_dir, output_path)
    presentation = Presentation(result)

    assert warnings == []
    assert len(presentation.slides) == 2
    assert all(
        len(slide.shapes) == 1 and slide.shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
        for slide in presentation.slides
    )
    report = json.loads(output_path.with_suffix(".image2_export_report.json").read_text(encoding="utf-8"))
    assert report["mode"] == "image2_image"
    assert report["export_strategy"] == "full_slide_image2_pages"


def test_svg_export_keeps_native_text_and_shapes(tmp_path: Path) -> None:
    svg_path = tmp_path / "01_editable.svg"
    svg_path.write_text(
        """
        <svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
          <rect x="0" y="0" width="1280" height="720" fill="#ffffff"/>
          <rect x="80" y="120" width="420" height="180" rx="18" fill="#fff7f2" stroke="#c43a12"/>
          <text x="110" y="210" font-size="38" fill="#12141a">可编辑 PPT 测试</text>
        </svg>
        """.strip(),
        encoding="utf-8",
    )
    output_path = tmp_path / "editable-output.pptx"

    create_pptx([svg_path], output_path, canvas_format="ppt169")
    presentation = Presentation(output_path)
    slide = presentation.slides[0]
    report = json.loads(output_path.with_suffix(".conversion_report.json").read_text(encoding="utf-8"))

    assert report[0]["mode"] == "native"
    assert any(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    assert any(getattr(shape, "has_text_frame", False) and "可编辑 PPT 测试" in shape.text for shape in slide.shapes)


def test_small_decks_use_proportional_llm_budgets() -> None:
    duration, chars = _scaled_script_targets(2, 3600, 8000)

    assert duration == 210
    assert chars == 800
    assert _script_token_budget(chars) == 4096
    assert _design_spec_token_budget(2) == 8192
    assert _design_spec_token_budget(35) == 24576
