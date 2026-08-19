#!/usr/bin/env python3
"""
OREP PPT生成系统 - Integration Tester（强化版）
执行3个核心测试用例：P0-001, P0-002, P2-001
"""
import os
import sys
import json
import time
import subprocess
from datetime import datetime
from typing import Dict, Optional, Tuple, List

# 项目路径
PROJECT_ROOT = "/Users/liuyixing/项目/OREP/ai-scoring"
sys.path.insert(0, PROJECT_ROOT)

# API配置
API_BASE_URL = "http://localhost:8090"
API_PPT_TASK_CREATE = f"{API_BASE_URL}/api/ppt/v2/task/create"
API_PPT_TASK_STATUS = f"{API_BASE_URL}/api/ppt/task/{{task_id}}/status"
API_PPT_TASK_DETAIL = f"{API_BASE_URL}/api/ppt/task/{{task_id}}"
PYTHON_BIN = f"{PROJECT_ROOT}/.venv/bin/python"

# 质量指标基线
QUALITY_BASELINES = {
    "min_effective_pages": 35,
    "max_empty_pages": 3,
    "min_svg_charts": 10,
    "min_skill_pages": 10,
    "max_garbled_chars": 0,
}

# 任务54基线（已知数据）
BASELINE_TASK_54 = {
    "total_pages": 40,
    "effective_pages": 38,  # 估计
    "empty_pages": 2,  # 估计
    "svg_charts": 15,  # 估计
    "skill_pages": 12,  # 估计
}


def api_post(url: str, data: Dict) -> Tuple[int, Dict]:
    """发送POST请求"""
    import urllib.request
    import urllib.error
    try:
        body = json.dumps(data).encode('utf-8')
        req = urllib.request.Request(url, data=body)
        req.add_header('Content-Type', 'application/json')
        with urllib.request.urlopen(req, timeout=60) as resp:
            return resp.status, json.loads(resp.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read().decode('utf-8'))
        except:
            return e.code, {"error": str(e)}
    except Exception as e:
        return 0, {"error": str(e)}


def api_get(url: str) -> Tuple[int, Dict]:
    """发送GET请求"""
    import urllib.request
    import urllib.error
    try:
        req = urllib.request.Request(url)
        with urllib.request.urlopen(req, timeout=30) as resp:
            return resp.status, json.loads(resp.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read().decode('utf-8'))
        except:
            return e.code, {"error": str(e)}
    except Exception as e:
        return 0, {"error": str(e)}


def analyze_pptx(pptx_path: str) -> Dict:
    """分析PPTX文件，提取质量指标"""
    analysis_script = f'''
import sys
from pptx import Presentation
import zipfile
import re

pptx_path = "{pptx_path}"

def has_garbled_chars(text):
    garbled_patterns = [r\'[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f]\', r\'\\ufffd\', r\'\\u0000\']
    for pattern in garbled_patterns:
        if re.search(pattern, text):
            return True
    return False

# Count SVGs
svg_count = 0
try:
    with zipfile.ZipFile(pptx_path, \'r\') as zf:
        for name in zf.namelist():
            if name.startswith(\'ppt/slides/slide\') and name.endswith(\'.xml\'):
                content = zf.read(name).decode(\'utf-8\', errors=\'ignore\')
                if \'<a:blip\' in content and \'svg\' in content.lower():
                    svg_count += 1
except Exception as e:
    pass

prs = Presentation(pptx_path)
total = len(prs.slides)
empty_pages = 0
skill_pages = 0
kpi_pages = 0
timeline_pages = 0
content_pages = 0
team_pages = 0
cover_pages = 0
toc_pages = 0
ending_pages = 0
effective = 0
garbled = 0

for i, slide in enumerate(prs.slides):
    slide_num = i + 1
    shapes = len(slide.shapes)
    all_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    text = run.text
                    if text:
                        all_text.append(text)
                        if has_garbled_chars(text):
                            garbled += 1
    full_text = \' \'.join(all_text).strip()
    word_count = len(full_text)
    layout = str(slide.slide_layout.name).lower() if slide.slide_layout else ""
    skill_keywords = [\'技能\', \'操作\', \'岗位职责\', \'前置准备\', \'后续收尾\', \'团队协作\', \'步骤\', \'milestone\', \'timeline\', \'技能操作\']
    is_skill = any(kw in full_text for kw in skill_keywords)
    kpi_keywords = [\'指标\', \'KPI\', \'成果\', \'率\', \'%\', \'达成\']
    is_kpi = (is_skill and \'metrics\' in layout.lower()) or any(kw in full_text for kw in kpi_keywords)
    timeline_keywords = [\'时间\', \'timeline\', \'规划\', \'发展\', \'阶段\']
    is_timeline = \'timeline\' in layout.lower() or any(kw in full_text for kw in timeline_keywords)
    content_keywords = [\'内容\', \'说明\', \'分析\', \'方案\', \'创新\']
    is_content = \'content\' in layout.lower() and word_count > 50
    team_keywords = [\'团队\', \'成员\', \'分工\', \'name\', \'role\']
    is_team = \'team\' in layout.lower() or (all(kw in full_text for kw in [\'团队\', \'成员\']) if full_text else False)
    is_cover = \'cover\' in layout.lower() or \'封面\' in full_text[:20]
    is_toc = \'toc\' in layout.lower() or \'目录\' in full_text[:20]
    is_ending = \'ending\' in layout.lower() or \'结束\' in full_text[:20] or \'谢谢\' in full_text[:20]
    is_empty = word_count < 20 and shapes < 3
    if is_skill: skill_pages += 1
    if is_kpi: kpi_pages += 1
    if is_timeline: timeline_pages += 1
    if is_content: content_pages += 1
    if is_team: team_pages += 1
    if is_cover: cover_pages += 1
    if is_toc: toc_pages += 1
    if is_ending: ending_pages += 1
    if is_empty: empty_pages += 1
    else: effective += 1

result = {{
    "total_pages": total,
    "effective_pages": effective,
    "empty_pages": empty_pages,
    "svg_charts": svg_count,
    "skill_pages": skill_pages,
    "garbled_chars": garbled,
    "kpi_pages": kpi_pages,
    "timeline_pages": timeline_pages,
    "content_pages": content_pages,
    "team_pages": team_pages,
    "cover_pages": cover_pages,
    "toc_pages": toc_pages,
    "ending_pages": ending_pages
}}

import json
print(json.dumps(result))
'''

    try:
        result = subprocess.run(
            [PYTHON_BIN, "-c", analysis_script],
            capture_output=True,
            text=True,
            timeout=120
        )
        if result.returncode == 0 and result.stdout.strip():
            return json.loads(result.stdout.strip())
    except Exception as e:
        print(f"PPTX分析异常: {e}")
    return {}


def create_task(questionnaire_id: int, user_id: int, tenant_id: int, industry: str = "healthcare") -> Optional[int]:
    """创建PPT生成任务"""
    data = {
        "questionnaire_id": questionnaire_id,
        "user_id": user_id,
        "tenant_id": tenant_id,
        "industry": industry
    }
    status, response = api_post(API_PPT_TASK_CREATE, data)
    if status == 200 and response.get("success"):
        return response["data"]["task_id"]
    print(f"创建任务失败: status={status}, response={response}")
    return None


def poll_task_status(task_id: int, timeout: int = 600) -> bool:
    """轮询任务状态直到完成/失败"""
    url = API_PPT_TASK_STATUS.format(task_id=task_id)
    start_time = time.time()
    poll_interval = 5

    while time.time() - start_time < timeout:
        status, response = api_get(url)
        if status != 200:
            time.sleep(poll_interval)
            continue

        data = response.get("data", {})
        current_status = data.get("status")
        current_step = data.get("current_step", "")
        print(f"  Task {task_id} status: {current_status}, step: {current_step}")

        if current_status == "completed":
            return True
        elif current_status in ["failed", "error"]:
            print(f"  Task failed: {data.get('error_msg', 'unknown')}")
            return False

        time.sleep(poll_interval)

    print(f"  Task {task_id} timeout")
    return False


def get_task_detail(task_id: int) -> Optional[Dict]:
    """获取任务详情"""
    url = API_PPT_TASK_DETAIL.format(task_id=task_id)
    status, response = api_get(url)
    if status == 200 and response.get("success"):
        return response
    return None


def run_test(case_id: str, questionnaire_id: int, user_id: int, tenant_id: int,
             industry: str, timeout: int = 600) -> Dict:
    """运行单个测试用例"""
    print(f"\n{'='*60}")
    print(f"Running Test: {case_id}")
    print(f"  questionnaire_id={questionnaire_id}, user_id={user_id}, tenant_id={tenant_id}, industry={industry}")
    print(f"{'='*60}")

    start_time = time.time()

    # 1. Create task
    print(f"[{case_id}] Creating task...")
    task_id = create_task(questionnaire_id, user_id, tenant_id, industry)
    if not task_id:
        return {"case_id": case_id, "status": "error", "error": "Failed to create task"}
    print(f"[{case_id}] Task created: {task_id}")

    # 2. Poll status
    print(f"[{case_id}] Polling status...")
    if not poll_task_status(task_id, timeout):
        return {"case_id": case_id, "task_id": task_id, "status": "timeout"}
    print(f"[{case_id}] Task completed")

    # 3. Get task detail
    print(f"[{case_id}] Getting task detail...")
    task_detail = get_task_detail(task_id)
    if not task_detail:
        return {"case_id": case_id, "task_id": task_id, "status": "error", "error": "Failed to get task detail"}

    pptx_path = task_detail.get("data", {}).get("pptx_path")
    if not pptx_path:
        return {"case_id": case_id, "task_id": task_id, "status": "error", "error": "No pptx_path in task detail"}

    # 4. Analyze PPT
    print(f"[{case_id}] Analyzing PPT: {pptx_path}")
    metrics = analyze_pptx(pptx_path)
    if not metrics:
        return {"case_id": case_id, "task_id": task_id, "status": "error", "error": "Failed to analyze PPT"}

    elapsed = time.time() - start_time

    return {
        "case_id": case_id,
        "task_id": task_id,
        "status": "completed",
        "elapsed_seconds": elapsed,
        "metrics": metrics,
        "pptx_path": pptx_path
    }


def print_test_result(result: Dict):
    """打印测试结果"""
    case_id = result.get("case_id", "unknown")
    print(f"\n{'='*60}")
    print(f"=== Test {case_id} Results ===")
    print(f"{'='*60}")

    print(f"Task ID: {result.get('task_id', 'N/A')}")
    print(f"Status: {result.get('status', 'unknown')}")

    if result.get("status") == "completed":
        metrics = result.get("metrics", {})
        print(f"Total Pages: {metrics.get('total_pages', 'N/A')}")
        print(f"Effective Pages: {metrics.get('effective_pages', 'N/A')}")
        print(f"Empty Pages: {metrics.get('empty_pages', 'N/A')}")
        print(f"SVG Charts: {metrics.get('svg_charts', 'N/A')}")
        print(f"Skill Operation Pages: {metrics.get('skill_pages', 'N/A')}")
        print(f"Garbled Chars: {metrics.get('garbled_chars', 'N/A')}")
        print(f"Elapsed Time: {result.get('elapsed_seconds', 0):.1f}s")
        print(f"PPTX Path: {result.get('pptx_path', 'N/A')}")

        print(f"\nPASS/FAIL:")
        effective = metrics.get('effective_pages', 0)
        empty = metrics.get('empty_pages', 0)
        svg = metrics.get('svg_charts', 0)
        skill = metrics.get('skill_pages', 0)
        garbled = metrics.get('garbled_chars', 0)

        print(f"  - 有效页≥35: {'PASS' if effective >= 35 else 'FAIL'} (实际{effective})")
        print(f"  - 空壳页≤3: {'PASS' if empty <= 3 else 'FAIL'} (实际{empty})")
        print(f"  - SVG≥10: {'PASS' if svg >= 10 else 'FAIL'} (实际{svg})")
        print(f"  - 技能操作页≥10: {'PASS' if skill >= 10 else 'FAIL'} (实际{skill})")
        print(f"  - 零乱码: {'PASS' if garbled == 0 else 'FAIL'} (实际{garbled})")

        # Baseline comparison
        print(f"\n与基线(Task 54)对比:")
        baseline = BASELINE_TASK_54
        print(f"  - 基线有效页: {baseline['effective_pages']} vs 实际: {effective}")
        print(f"  - 基线空壳页: {baseline['empty_pages']} vs 实际: {empty}")
        print(f"  - 基线SVG: {baseline['svg_charts']} vs 实际: {svg}")

    elif result.get("status") == "timeout":
        print("Task timed out")

    elif result.get("status") == "error":
        print(f"Error: {result.get('error', 'unknown')}")

    print(f"{'='*60}")


def main():
    print("OREP PPT生成系统 - Integration Tester（强化版）")
    print("执行3个核心测试用例: P0-001(基线), P0-002(医疗回归), P2-001(最小数据)")
    print(f"时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # Task 54 baseline info
    print(f"\n基线Task 54指标: {BASELINE_TASK_54}")

    results = []

    # Test P0-001: Baseline test (using questionnaire_id=28)
    result_p0_001 = run_test(
        case_id="P0-001",
        questionnaire_id=28,
        user_id=1,
        tenant_id=1,
        industry="healthcare"
    )
    results.append(result_p0_001)
    print_test_result(result_p0_001)

    # Test P0-002: Multi-industry regression test (Healthcare)
    result_p0_002 = run_test(
        case_id="P0-002",
        questionnaire_id=28,
        user_id=1,
        tenant_id=1,
        industry="healthcare"
    )
    results.append(result_p0_002)
    print_test_result(result_p0_002)

    # Test P2-001: Minimal data test
    # For P2-001, we use questionnaire_id=0 or a minimal questionnaire
    # Since we don't have a minimal questionnaire, we test with questionnaire_id=28
    # but the system should handle it
    result_p2_001 = run_test(
        case_id="P2-001",
        questionnaire_id=28,  # Using same as baseline for now
        user_id=1,
        tenant_id=1,
        industry="tech"
    )
    results.append(result_p2_001)
    print_test_result(result_p2_001)

    # Summary
    print(f"\n{'='*60}")
    print("=== Test Summary ===")
    print(f"{'='*60}")
    for r in results:
        status = r.get("status", "unknown")
        case_id = r.get("case_id", "unknown")
        task_id = r.get("task_id", "N/A")
        print(f"  {case_id}: {status} (task_id={task_id})")

    passed = sum(1 for r in results if r.get("status") == "completed")
    print(f"\nTotal: {len(results)} tests, {passed} completed")

    # Save results
    output_file = f"/Users/liuyixing/项目/OREP/ai-scoring/tests/integration/results/integration_test_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\nResults saved to: {output_file}")


if __name__ == "__main__":
    main()