#!/usr/bin/env python3
"""Quick PPT analysis to verify content"""
from pptx import Presentation
import zipfile
import re
import json
import sys

def has_garbled_chars(text):
    garbled_patterns = [r'[\x00-\x08\x0b\x0c\x0e-\x1f]', r'\ufffd', r'\u0000']
    for pattern in garbled_patterns:
        if re.search(pattern, text):
            return True
    return False

def analyze_pptx(pptx_path):
    # Count SVGs
    svg_count = 0
    try:
        with zipfile.ZipFile(pptx_path, 'r') as zf:
            for name in zf.namelist():
                if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
                    content = zf.read(name).decode('utf-8', errors='ignore')
                    if '<a:blip' in content and 'svg' in content.lower():
                        svg_count += 1
    except Exception as e:
        print(f"SVG error: {e}", file=sys.stderr)

    prs = Presentation(pptx_path)
    total = len(prs.slides)

    empty_pages = 0
    skill_pages = 0
    kpi_pages = 0
    timeline_pages = 0
    content_pages = 0
    team_pages = 0
    cover_pages = 0
    toc_pages = 0
    ending_pages = 0
    effective = 0
    garbled = 0

    page_details = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        shapes = len(slide.shapes)

        all_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        text = run.text
                        if text:
                            all_text.append(text)
                            if has_garbled_chars(text):
                                garbled += 1

        full_text = ' '.join(all_text).strip()
        word_count = len(full_text)

        layout = str(slide.slide_layout.name).lower() if slide.slide_layout else ""

        skill_keywords = ['技能', '操作', '岗位职责', '前置准备', '后续收尾', '团队协作', '步骤', 'milestone', 'timeline', '技能操作']
        is_skill = any(kw in full_text for kw in skill_keywords)

        kpi_keywords = ['指标', 'KPI', '成果', '率', '%', '达成']
        is_kpi = (is_skill and 'metrics' in layout.lower()) or any(kw in full_text for kw in kpi_keywords)

        timeline_keywords = ['时间', 'timeline', '规划', '发展', '阶段']
        is_timeline = 'timeline' in layout.lower() or any(kw in full_text for kw in timeline_keywords)

        content_keywords = ['内容', '说明', '分析', '方案', '创新']
        is_content = 'content' in layout.lower() and word_count > 50

        team_keywords = ['团队', '成员', '分工', 'name', 'role']
        is_team = 'team' in layout.lower() or (all(kw in full_text for kw in ['团队', '成员']) if full_text else False)

        is_cover = 'cover' in layout.lower() or '封面' in full_text[:20]
        is_toc = 'toc' in layout.lower() or '目录' in full_text[:20]
        is_ending = 'ending' in layout.lower() or '结束' in full_text[:20] or '谢谢' in full_text[:20]

        # Be more lenient: only consider empty if truly no content
        is_empty = word_count < 10 and shapes < 2

        if is_skill: skill_pages += 1
        if is_kpi: kpi_pages += 1
        if is_timeline: timeline_pages += 1
        if is_content: content_pages += 1
        if is_team: team_pages += 1
        if is_cover: cover_pages += 1
        if is_toc: toc_pages += 1
        if is_ending: ending_pages += 1
        if is_empty: empty_pages += 1
        else: effective += 1

        page_details.append({
            "slide_num": slide_num,
            "shapes": shapes,
            "word_count": word_count,
            "is_empty": is_empty,
            "is_skill": is_skill,
            "is_kpi": is_kpi,
            "is_timeline": is_timeline,
            "layout": layout,
            "preview": full_text[:100] if full_text else "(empty)"
        })

    result = {
        "total_pages": total,
        "effective_pages": effective,
        "empty_pages": empty_pages,
        "svg_charts": svg_count,
        "skill_pages": skill_pages,
        "garbled_chars": garbled,
        "kpi_pages": kpi_pages,
        "timeline_pages": timeline_pages,
        "content_pages": content_pages,
        "team_pages": team_pages,
        "cover_pages": cover_pages,
        "toc_pages": toc_pages,
        "ending_pages": ending_pages,
        "page_details": page_details[:5]  # Only first 5 for brevity
    }

    return result

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python quick_ppt_check.py <pptx_path>")
        sys.exit(1)

    pptx_path = sys.argv[1]
    result = analyze_pptx(pptx_path)
    print(json.dumps(result, ensure_ascii=False, indent=2))