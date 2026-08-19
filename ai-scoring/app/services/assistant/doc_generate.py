"""将助手 Markdown/纯文本渲染为 DOCX 字节。"""
from __future__ import annotations

import io
import re
from typing import Iterable


def _iter_blocks(text: str) -> Iterable[tuple[str, str]]:
    """粗分块：heading / list / code / para。"""
    lines = (text or "").replace("\r\n", "\n").split("\n")
    buf: list[str] = []
    in_code = False
    code_lang = ""

    def flush_para():
        nonlocal buf
        if not buf:
            return
        yield ("para", "\n".join(buf).strip())
        buf = []

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                yield ("code", "\n".join(buf))
                buf = []
                in_code = False
                code_lang = ""
            else:
                if buf:
                    yield from flush_para()
                in_code = True
                code_lang = line.strip()[3:].strip()
                buf = []
            i += 1
            continue
        if in_code:
            buf.append(line)
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if m:
            if buf:
                yield from flush_para()
            level = len(m.group(1))
            yield (f"h{level}", m.group(2).strip())
            i += 1
            continue

        if re.match(r"^[-*]\s+", line.strip()) or re.match(r"^\d+\.\s+", line.strip()):
            if buf and not all(re.match(r"^([-*]|\d+\.)\s+", x.strip()) for x in buf if x.strip()):
                yield from flush_para()
            buf.append(line)
            i += 1
            # continue collecting list lines
            while i < len(lines) and (
                re.match(r"^[-*]\s+", lines[i].strip())
                or re.match(r"^\d+\.\s+", lines[i].strip())
                or (lines[i].startswith("  ") and lines[i].strip())
            ):
                buf.append(lines[i])
                i += 1
            yield ("list", "\n".join(buf))
            buf = []
            continue

        if not line.strip():
            if buf:
                yield from flush_para()
            i += 1
            continue

        buf.append(line)
        i += 1

    if in_code and buf:
        yield ("code", "\n".join(buf))
    elif buf:
        yield from flush_para()


def markdown_to_docx_bytes(text: str, title: str | None = None) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_LINE_SPACING
    except Exception as e:
        raise RuntimeError(f"python-docx 不可用: {e}") from e

    doc = Document()
    if title:
        doc.add_heading(title[:120], level=1)

    for kind, content in _iter_blocks(text):
        if not content and kind != "code":
            continue
        if kind == "h1":
            doc.add_heading(content, level=1)
        elif kind == "h2":
            doc.add_heading(content, level=2)
        elif kind == "h3":
            doc.add_heading(content, level=3)
        elif kind == "list":
            for line in content.split("\n"):
                s = line.strip()
                if not s:
                    continue
                s = re.sub(r"^[-*]\s+", "", s)
                s = re.sub(r"^\d+\.\s+", "", s)
                # strip simple markdown bold
                s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
                doc.add_paragraph(s, style="List Bullet")
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        else:
            # para: light markdown strip
            s = re.sub(r"\*\*(.+?)\*\*", r"\1", content)
            s = re.sub(r"`([^`]+)`", r"\1", s)
            doc.add_paragraph(s)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _slides_from_markdown(text: str) -> list[tuple[str, list[str]]]:
    """Markdown → [(title, bullets)]，用于简易 PPT。

    约定（与 Brain PPT 契约对齐）：
    - `#` 总标题：只作封面候选，不单独成空内容页
    - `##` / `###`：一页
    - 跳过 `---`、客服句、过长元说明
    """
    slides: list[tuple[str, list[str]]] = []
    deck_title: str | None = None
    title = ""
    bullets: list[str] = []

    def flush():
        nonlocal title, bullets
        t = (title or "").strip()
        bs = [b for b in bullets if b and b not in ("---", "***")][:8]
        # 丢掉完全空页
        if not t and not bs:
            bullets = []
            return
        # 丢掉纯元信息页
        if t and re.search(r"结构与内容|制作与演示|制作建议|核心建议", t):
            title = ""
            bullets = []
            return
        if not t and not bs:
            return
        if t or bs:
            slides.append((t or "内容", bs[:6] if bs else ["（要点见讲稿）"]))
        title = ""
        bullets = []

    for line in (text or "").replace("\r\n", "\n").split("\n"):
        s = line.strip()
        if not s or s.startswith("```"):
            continue
        if s in ("---", "***", "___"):
            continue
        if re.match(r"^(如需我|请告知|如果你愿意)", s):
            continue
        m = re.match(r"^(#{1,3})\s+(.+)$", s)
        if m:
            level = len(m.group(1))
            heading = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)).strip()
            heading = re.sub(r"^第\s*\d+\s*[-~～到至]\s*\d+\s*页[：:：\s]*", "核心功能 · ", heading)
            if level == 1 and deck_title is None:
                deck_title = heading
                # 不 flush 成内容页；封面用 deck_title
                continue
            flush()
            title = heading
            continue
        if re.match(r"^[-*]\s+", s) or re.match(r"^\d+[\.、]\s*", s):
            s = re.sub(r"^[-*]\s+", "", s)
            s = re.sub(r"^\d+[\.、]\s*", "", s)
            s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
            s = re.sub(r"`([^`]+)`", r"\1", s)
            # 表格残留
            if s.startswith("|"):
                cells = [c.strip() for c in s.strip("|").split("|") if c.strip()]
                s = " · ".join(cells[:4])
            if s and len(s) > 1:
                bullets.append(s[:100])
            continue
        # 普通段落：过长元描述跳过，短句作要点
        s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
        s = s.lstrip("> ").strip()
        if re.search(r"建议总页数|本PPT结构|依据《", s) and len(s) > 40:
            continue
        if 2 < len(s) <= 80:
            bullets.append(s[:100])
            if len(bullets) >= 8:
                flush()
                title = "续"
        elif len(s) > 80:
            # 长段压成一句
            bullets.append(s[:90] + "…")
    flush()
    if not slides:
        slides = [(deck_title or "竞赛助手提纲", ["（暂无结构化要点，请按「每页一个 ## 标题」重试）"])]
    return slides[:18]


def _new_prs():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def deck_schema_to_pptx_bytes(schema: dict) -> bytes:
    """结构化 deck JSON → PPTX（P2 主路径）。"""
    try:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
    except Exception as e:
        raise RuntimeError(f"python-pptx 不可用: {e}") from e

    prs = _new_prs()
    blank = prs.slide_layouts[6]
    deck_title = (schema.get("title") or "路演PPT")[:80]
    slides = schema.get("slides") or []

    ORANGE = RGBColor(0xC4, 0x3A, 0x12)
    INK = RGBColor(0x12, 0x14, 0x1A)
    MUTED = RGBColor(0x6B, 0x72, 0x80)
    BODY = RGBColor(0x2C, 0x30, 0x38)

    def add_title_bar(slide, text: str, y=0.4):
        tbox = slide.shapes.add_textbox(Inches(0.7), Inches(y), Inches(12), Inches(0.9))
        tf = tbox.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = (text or "")[:60]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = INK
        # 橙色短条
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(y + 0.85), Inches(1.2), Inches(0.06)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = ORANGE
        bar.line.fill.background()

    for s in slides:
        layout = (s.get("layout") or "bullets").lower()
        title = (s.get("title") or deck_title)[:60]
        subtitle = (s.get("subtitle") or "")[:100]
        bullets = list(s.get("bullets") or [])[:6]
        rows = list(s.get("rows") or [])[:8]

        slide = prs.slides.add_slide(blank)

        if layout == "cover":
            # 左侧色块
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(7.5)
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = ORANGE
            accent.line.fill.background()

            box = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.2), Inches(2.8))
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title[:80]
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = ORANGE
            p.alignment = PP_ALIGN.LEFT
            if subtitle:
                sp = tf.add_paragraph()
                sp.text = subtitle
                sp.font.size = Pt(16)
                sp.font.color.rgb = MUTED
                sp.space_before = Pt(14)
            foot = tf.add_paragraph()
            foot.text = "竞赛大脑 · 竞赛助手"
            foot.font.size = Pt(12)
            foot.font.color.rgb = MUTED
            foot.space_before = Pt(28)
            continue

        if layout == "closing":
            add_title_bar(slide, title or "致谢")
            box = slide.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.2), Inches(4.5))
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            items = bullets or ["感谢各位评委老师", "欢迎提问交流"]
            for i, b in enumerate(items):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = b
                para.font.size = Pt(22)
                para.font.color.rgb = BODY
                para.space_after = Pt(14)
                para.alignment = PP_ALIGN.LEFT
            if subtitle:
                sp = tf.add_paragraph()
                sp.text = subtitle
                sp.font.size = Pt(14)
                sp.font.color.rgb = MUTED
            continue

        # bullets / compare / default
        add_title_bar(slide, title)
        if subtitle:
            sbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.35), Inches(12), Inches(0.4))
            stf = sbox.text_frame
            stf.clear()
            sp = stf.paragraphs[0]
            sp.text = subtitle
            sp.font.size = Pt(14)
            sp.font.color.rgb = MUTED

        top = 1.7 if subtitle else 1.5

        if layout == "compare" and rows:
            # 简易表格区域：每行一条
            box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(5.0))
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            # header if first row looks like header
            for i, row in enumerate(rows):
                if not isinstance(row, (list, tuple)):
                    continue
                cells = [str(c) for c in row if str(c).strip()]
                line = "  |  ".join(cells[:4])
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = ("▸ " if i else "● ") + line
                para.font.size = Pt(16 if i else 17)
                para.font.bold = i == 0
                para.font.color.rgb = INK if i == 0 else BODY
                para.space_after = Pt(12)
            if bullets:
                for b in bullets[:3]:
                    para = tf.add_paragraph()
                    para.text = f"• {b}"
                    para.font.size = Pt(15)
                    para.font.color.rgb = MUTED
        else:
            box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(5.0))
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            if not bullets:
                bullets = ["（本页要点待补充）"]
            for i, b in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = f"• {b}"
                para.font.size = Pt(18)
                para.font.color.rgb = BODY
                para.level = 0
                para.space_after = Pt(12)

    if len(prs.slides) == 0:
        # 兜底一页
        slide = prs.slides.add_slide(blank)
        add_title_bar(slide, deck_title)

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def markdown_to_pptx_bytes(text: str, title: str = "竞赛助手提纲") -> bytes:
    """Markdown → PPTX（兼容旧路径；内部转 deck schema）。"""
    from app.services.assistant.artifact_schema import deck_from_markdown

    schema = deck_from_markdown(text, title=title)
    return deck_schema_to_pptx_bytes(schema)


def plan_schema_to_xlsx_bytes(schema: dict) -> bytes:
    """plan JSON → xlsx。"""
    try:
        import xlsxwriter
    except Exception as e:
        raise RuntimeError(f"xlsxwriter 不可用: {e}") from e

    bio = io.BytesIO()
    title = (schema.get("title") or "计划")[:31] or "计划"
    wb = xlsxwriter.Workbook(bio, {"in_memory": True})
    ws = wb.add_worksheet(title)
    header = wb.add_format({"bold": True, "bg_color": "#FFF3ED", "border": 1})
    cell = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
    title_fmt = wb.add_format({"bold": True, "font_size": 14})

    ws.set_column(0, 0, 8)
    ws.set_column(1, 1, 28)
    ws.set_column(2, 2, 40)
    ws.set_column(3, 3, 14)
    ws.set_column(4, 4, 14)

    ws.write(0, 0, schema.get("title") or "训练计划", title_fmt)
    headers = ["序号", "事项", "说明/目标", "负责人", "时间"]
    for col, h in enumerate(headers):
        ws.write(1, col, h, header)

    for i, r in enumerate(schema.get("rows") or [], 1):
        row = i + 1
        ws.write(row, 0, i, cell)
        ws.write(row, 1, r.get("item") or "", cell)
        ws.write(row, 2, r.get("detail") or "", cell)
        ws.write(row, 3, r.get("owner") or "", cell)
        ws.write(row, 4, r.get("due") or "", cell)

    wb.close()
    return bio.getvalue()


def doc_schema_to_docx_bytes(schema: dict) -> bytes:
    """doc JSON → docx。"""
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as e:
        raise RuntimeError(f"python-docx 不可用: {e}") from e

    doc = Document()
    title = schema.get("title") or "竞赛助手文稿"
    doc.add_heading(str(title)[:120], level=1)
    for sec in schema.get("sections") or []:
        h = (sec.get("heading") or "").strip()
        level = int(sec.get("level") or 2)
        if h and h != title:
            doc.add_heading(h[:120], level=min(max(level, 1), 3))
        for p in sec.get("paras") or []:
            if p:
                doc.add_paragraph(str(p))
        for b in sec.get("bullets") or []:
            if b:
                doc.add_paragraph(str(b), style="List Bullet")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def markdown_to_xlsx_bytes(text: str, sheet_title: str = "计划") -> bytes:
    """把 Markdown/列表文本整理成简易 Excel。"""
    try:
        import xlsxwriter
    except Exception as e:
        raise RuntimeError(f"xlsxwriter 不可用: {e}") from e

    bio = io.BytesIO()
    wb = xlsxwriter.Workbook(bio, {"in_memory": True})
    ws = wb.add_worksheet((sheet_title or "计划")[:31] or "计划")
    header = wb.add_format({"bold": True, "bg_color": "#FFF3ED", "border": 1})
    cell = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
    title_fmt = wb.add_format({"bold": True, "font_size": 14})

    ws.set_column(0, 0, 8)
    ws.set_column(1, 1, 28)
    ws.set_column(2, 2, 36)
    ws.set_column(3, 3, 20)

    ws.write(0, 0, "竞赛助手导出", title_fmt)
    ws.write(1, 0, "序号", header)
    ws.write(1, 1, "事项", header)
    ws.write(1, 2, "说明/目标", header)
    ws.write(1, 3, "备注", header)

    rows: list[tuple[str, str]] = []
    for line in (text or "").replace("\r\n", "\n").split("\n"):
        s = line.strip()
        if not s:
            continue
        if s.startswith("#"):
            s = re.sub(r"^#+\s*", "", s)
            rows.append((s, "章节/标题"))
            continue
        if re.match(r"^[-*]\s+", s) or re.match(r"^\d+[\.、]\s*", s):
            s = re.sub(r"^[-*]\s+", "", s)
            s = re.sub(r"^\d+[\.、]\s*", "", s)
            s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
            # 尝试拆「事项：说明」
            if "：" in s:
                a, b = s.split("：", 1)
                rows.append((a.strip(), b.strip()))
            elif ":" in s:
                a, b = s.split(":", 1)
                rows.append((a.strip(), b.strip()))
            else:
                rows.append((s, ""))
            continue
        # 普通段落
        s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
        if len(s) > 4:
            rows.append((s[:40] + ("…" if len(s) > 40 else ""), s))

    if not rows:
        rows = [("（暂无条目）", "可根据对话内容继续完善")]

    for i, (item, desc) in enumerate(rows[:200], 1):
        r = i + 1
        ws.write(r, 0, i, cell)
        ws.write(r, 1, item, cell)
        ws.write(r, 2, desc, cell)
        ws.write(r, 3, "", cell)

    wb.close()
    return bio.getvalue()
